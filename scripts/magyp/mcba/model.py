"""Parser XLSX MCBA y modelo market_price_observation v1.

No conserva formularios HTML con tokens GeneXus ni inventa volumen.
"""
from __future__ import annotations

import io
import json
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from statistics import median

from openpyxl import load_workbook

from scripts.magyp.common.platform import PipelineError, RawManifest, stable_id

PARSER_VERSION = "mcba-xlsx-1.0.0"
SCHEMA_VERSION = "market_price_observation-v1"
ENDPOINT = "https://ssma.magyp.gob.ar/frutas.precios.aspx"
CONTEXT_URL = "https://ssma.magyp.gob.ar/frutas.preciospromediof.aspx"
HEADERS = ["Fecha", "Tipo", "Especie", "Variedad", "Procedencia", "Envase", "Calidad",
           "Tamaño", "Grado", "Kg", "Promedio x Kg."]
FIELDS = ["observation_id", "source", "source_family", "market", "market_normalized",
          "observation_date", "observation_datetime", "period", "frequency", "date_precision",
          "product", "product_normalized", "variety", "variety_normalized", "origin",
          "origin_province", "origin_locality", "package", "package_normalized", "quality",
          "quality_detail", "type", "size", "grade", "currency", "price_unit", "price_min",
          "price_max", "price_average", "price_modal", "price", "volume", "volume_unit",
          "source_record_id", "capture_id", "capture_timestamp", "raw_sha256", "source_url",
          "schema_version", "parser_version", "parse_warnings", "quality_flags", "original_dimensions",
          "price_raw", "kg_raw", "source_row_number", "record_kind", "currency_basis", "unit_basis",
          "dimension_key", "record_fingerprint"]


def text(value):
    if value is None or not str(value).strip():
        return None
    return str(value).strip()


def canonical(value):
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).upper().split()) or None


def number(value):
    if value is None or value == "":
        return None
    try:
        s = str(value).strip()
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        d = Decimal(s)
        return float(d) if d.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value), fmt).date().isoformat()
        except ValueError:
            pass
    return None


def parse_export(payload: bytes, requested_date: str) -> list[dict]:
    """Preserva filas con valores inválidos. Rechaza schema, vacío o ventana distinta."""
    if not payload.startswith(b"PK"):
        raise PipelineError("Respuesta no XLSX (posible HTML/error de servicio)")
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as z:
            if sum(x.file_size for x in z.infolist()) > 20_000_000 or len(z.infolist()) > 200:
                raise PipelineError("XLSX excede límite descomprimido")
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        sheet = book.active
        if sheet.max_row > 10005 or sheet.max_column != 11:
            raise PipelineError("Dimensiones inesperadas del XLSX")
        rows = list(sheet.iter_rows(values_only=True))
        book.close()
    except (zipfile.BadZipFile, OSError, KeyError, ValueError) as e:
        raise PipelineError("XLSX inválido") from e
    header = next((i for i, r in enumerate(rows[:10]) if list(r) == HEADERS), None)
    if header is None:
        raise PipelineError("Schema MCBA cambiado: encabezados no reconocidos")
    records = []
    for i, cells in enumerate(rows[header + 1:], header + 2):
        if all(v is None or v == "" for v in cells):
            continue
        record = dict(zip(HEADERS, cells))
        parsed = iso_date(record["Fecha"])
        if parsed and parsed != requested_date:
            raise PipelineError("Exportación fuera del día solicitado: captura rechazada")
        record["_row"] = i
        records.append(record)
    if not records:
        raise PipelineError("Exportación vacía; ausencia de publicación no demostrada")
    if not any(iso_date(x["Fecha"]) == requested_date for x in records):
        raise PipelineError("Ninguna fecha del export confirma el día solicitado")
    return records


def normalize(records: list[dict], manifest: RawManifest) -> list[dict]:
    results = []
    counts = Counter()
    for r in records:
        row = dict.fromkeys(FIELDS)
        d = iso_date(r["Fecha"])
        price = number(r["Promedio x Kg."])
        raw = {k: v.isoformat() if isinstance(v, (datetime, date)) else v for k, v in r.items() if k != "_row"}
        warnings = ["kg_semantics_unverified", "official_record_id_not_exported"]
        if any(isinstance(v, str) and ("\u00a5" in v or "\ufffd" in v) for v in raw.values()):
            warnings.append("suspicious_source_character_preserved")
        if d is None:
            warnings.append("date_parse_failed")
        if price is None and text(r["Promedio x Kg."]) is not None:
            warnings.append("price_parse_failed")
        if manifest.currency:
            warnings.append("currency_from_official_context")
            if canonical(r["Tipo"]) == "HORTALIZAS":
                warnings.append("currency_context_transfer_requires_confirmation")
        row.update(source="mcba", source_family="market_price_observation",
                   market="Mercado Central de Buenos Aires", market_normalized="MCBA",
                   observation_date=d, period=d[:7] if d else None,
                   frequency="daily", date_precision="day", product=text(r["Especie"]),
                   product_normalized=canonical(r["Especie"]), variety=text(r["Variedad"]),
                   variety_normalized=canonical(r["Variedad"]), origin=text(r["Procedencia"]),
                   package=text(r["Envase"]), package_normalized=canonical(r["Envase"]),
                   quality=text(r["Calidad"]), type=text(r["Tipo"]), size=text(r["Tamaño"]),
                   grade=text(r["Grado"]), currency=manifest.currency, price_unit=manifest.price_unit,
                   price_average=price, price=price, capture_id=manifest.capture_id,
                   capture_timestamp=manifest.captured_at_utc, raw_sha256=manifest.sha256,
                   source_url=manifest.endpoint, schema_version=SCHEMA_VERSION, parser_version=PARSER_VERSION,
                   parse_warnings=warnings, original_dimensions=raw, price_raw=raw["Promedio x Kg."],
                   kg_raw=raw["Kg"], source_row_number=r["_row"],
                   record_kind="species_summary" if canonical(r["Variedad"]) == "PROM.ESP." else "detail",
                   currency_basis="official_context" if manifest.currency else "not_reported",
                   unit_basis="export_header_and_currency_context" if manifest.price_unit else "not_reported")
        dimensions = [row[k] for k in ("source", "market_normalized", "observation_date", "type", "product",
                      "variety", "origin", "package", "quality", "size", "grade", "price_unit", "currency")]
        row["dimension_key"] = stable_id(dimensions)
        row["record_fingerprint"] = stable_id(raw)
        occurrence = counts[row["record_fingerprint"]]
        counts[row["record_fingerprint"]] += 1
        row["observation_id"] = stable_id(["mcba", row["record_fingerprint"], occurrence])
        row["quality_flags"] = quality_flags(row)
        results.append(row)
    candidates = Counter(x["dimension_key"] for x in results)
    for row in results:
        if candidates[row["dimension_key"]] > 1:
            row["quality_flags"].append("duplicate_candidate")
    flag_extremes(results)
    return results


def quality_flags(row):
    flags = []
    price = row["price"]
    if price is None:
        flags.append("price_missing")
    elif price == 0:
        flags.append("price_zero")
    elif price < 0:
        flags.append("price_negative")
    for name, field in (("currency_missing", "currency"), ("unit_missing", "price_unit"),
                        ("product_missing", "product"), ("date_missing", "observation_date"),
                        ("origin_missing", "origin")):
        if not row[field]:
            flags.append(name)
    return flags


def flag_extremes(rows):
    # Diagnóstico relativo, nunca descarte automático: grupo >=5, precio >10x mediana.
    groups = defaultdict(list)
    for r in rows:
        if r["price"] is not None and r["price"] > 0:
            key = (r["observation_date"], r["product_normalized"], r["record_kind"], r["currency"], r["price_unit"])
            groups[key].append(r)
    for group in groups.values():
        if len(group) >= 5:
            center = median(r["price"] for r in group)
            for row in group:
                if row["price"] > 10 * center or row["price"] < center / 10:
                    row["quality_flags"].append("extreme_value_candidate")


def validate_canonical(rows: list[dict], analytical=False) -> None:
    if not rows:
        raise PipelineError("Dataset vacío")
    required = set(FIELDS) | ({"valid_for_price_series", "reason"} if analytical else set())
    seen = set()
    for row in rows:
        if set(row) != required or row["schema_version"] != SCHEMA_VERSION or row["source"] != "mcba":
            raise PipelineError("Contrato canónico inválido")
        key = (row["capture_id"], row["observation_id"])
        if key in seen:
            raise PipelineError("ID duplicado dentro de captura")
        seen.add(key)
        if len(row["raw_sha256"] or "") != 64 or not row["capture_timestamp"]:
            raise PipelineError("Trazabilidad incompleta")
        if row["observation_date"] and iso_date(row["observation_date"]) != row["observation_date"]:
            raise PipelineError("Fecha canónica inválida")
        if number(row["price"]) != row["price"] or not isinstance(row["quality_flags"], list):
            raise PipelineError("Precio/flags inválidos")
        if not set(quality_flags(row)).issubset(row["quality_flags"]):
            raise PipelineError("Flags inconsistentes")


def build_analytical(rows):
    validate_canonical(rows)
    # Revisión por publicación diaria completa, nunca rellenar dimensiones con capturas antiguas.
    latest = {}
    for row in rows:
        key = row["observation_date"] or ("invalid_date", row["capture_id"])
        candidate = (row["capture_timestamp"], row["capture_id"])
        latest[key] = max(latest.get(key, candidate), candidate)
    result = []
    blockers = {"price_missing", "price_zero", "price_negative", "currency_missing", "unit_missing",
                "product_missing", "date_missing", "duplicate_candidate"}
    for row in rows:
        key = row["observation_date"] or ("invalid_date", row["capture_id"])
        if (row["capture_timestamp"], row["capture_id"]) != latest[key]:
            continue
        out = dict(row)
        reasons = sorted(blockers.intersection(row["quality_flags"]))
        out.update(valid_for_price_series=not reasons, reason=";".join(reasons) or "valid")
        result.append(out)
    validate_canonical(result, analytical=True)
    return result


def dashboard_rows(rows):
    validate_canonical(rows, analytical=True)
    # Todas las dimensiones, incluyendo origen/calidad/tamaño/grado y Prom.Esp., siguen separadas.
    dimensions = ["source", "market_normalized", "type", "product_normalized", "variety_normalized",
                  "origin", "package_normalized", "quality", "size", "grade", "currency", "price_unit", "record_kind"]
    daily = []
    for r in rows:
        if r["valid_for_price_series"]:
            daily.append({**{k: r[k] for k in dimensions}, "date": r["observation_date"],
                          "price": r["price"], "observation_id": r["observation_id"],
                          "source_url": r["source_url"], "updated_at": r["capture_timestamp"],
                          "quality_flags": r["quality_flags"]})
    if not daily:
        raise PipelineError("No hay precios válidos: última salida conservada")
    daily.sort(key=lambda r: (r["date"], str(tuple(r[k] for k in dimensions)), r["observation_id"]))
    groups = defaultdict(list)
    last = {}
    for row in daily:
        key = tuple(row[k] for k in dimensions)
        groups[(row["date"][:7], key)].append(row)
        if key not in last or (row["date"], row["updated_at"], row["observation_id"]) > (
                last[key]["date"], last[key]["updated_at"], last[key]["observation_id"]):
            last[key] = row
    monthly = []
    for (period, key), group in sorted(groups.items(), key=lambda x: str(x[0])):
        monthly.append({**dict(zip(dimensions, key)), "period": period,
                        "price_median": median(r["price"] for r in group), "observation_count": len(group),
                        "observed_days": len({r["date"] for r in group}),
                        "coverage_status": "partial_unverified", "method": "median_observed_daily_prices",
                        "source_url": ENDPOINT, "updated_at": max(r["updated_at"] for r in group)})
    flags = Counter(f for r in rows for f in r["quality_flags"])
    summary = [{"source": "mcba", "rows_analytical": len(rows), "valid_rows": len(daily),
                "date_min": min(r["date"] for r in daily), "date_max": max(r["date"] for r in daily),
                "observed_days": len({r["date"] for r in daily}), "detail_rows": sum(r["record_kind"] == "detail" for r in rows),
                "summary_rows": sum(r["record_kind"] == "species_summary" for r in rows),
                "quality_flags": dict(flags), "source_url": ENDPOINT,
                "updated_at": max(r["capture_timestamp"] for r in rows)}]
    return {"DAILY": daily, "MONTHLY": monthly, "LATEST": list(last.values()), "SUMMARY": summary}
