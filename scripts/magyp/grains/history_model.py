"""Contratos históricos por familia. No relaja los parsers de actualización diaria."""
from __future__ import annotations
import csv
import io
import json
import re
import unicodedata
from datetime import date, datetime
from openpyxl import load_workbook
from scripts.magyp.grains import references as r

VERSION = 'magyp-history-v1'


class NoPublication(r.Error):
    status='NO_PUBLICATION'


class Unavailable(r.Error):
    status='SOURCE_UNAVAILABLE'


def internal_annual(obj,m,as_of):
    if ('PRECIOS INTERNOS DE LOS PRINCIPALES GRANOS ( $/TN.)' not in label(obj['text'])
            and 'Pesos Argentinos por Toneladas - $/Ton.' not in obj['text']):
        raise r.Error('Definición/unidad interna anual desconocida')
    named=re.search(r'_(20\d\d)\.php$',m['source_url'])
    if not named:raise r.Error('Año interno no identificado')
    block=None;products=None;markets=None;year=None;seen=set();blocks=set()
    for index,cells in enumerate(obj['rows']):
        if len(cells)==4 and cells[1:]==['TRIGO','MAIZ','SORGO']:
            block=0;products=None;markets=None;year=None;blocks.add(block);continue
        if len(cells)==4 and cells[1:3]==['SOJA','GIRASOL'] and cells[3] in ('CEBADA FORRAJERA','LINO'):
            block=1;third=cells[3];products=None;markets=None;year=None;blocks.add(block);continue
        if block is None:continue
        if cells and cells[0]=='DARSENA':
            if block==0:
                wheat=['DARSENA','ROSARIO','B.BLANCA']+(['CORDOBA'] if int(named[1])>=2020 else [])+['QUEQUEN']
                expected=wheat+['DARSENA','ROSARIO','B.BLANCA','QUEQUEN','DARSENA','ROSARIO','B.BLANCA']
                products=['TRIGO']*len(wheat)+['MAIZ']*4+['SORGO']*3
            else:
                tail=['B.BLANCA','QUEQUEN'] if third=='CEBADA FORRAJERA' else ['DARSENA','ROSARIO','QUEQUEN']
                expected=['DARSENA','ROSARIO','B.BLANCA','QUEQUEN']*2+tail
                products=['SOJA']*4+['GIRASOL']*4+[third]*len(tail)
            if cells!=expected:raise r.Error('Plazas internas anuales cambiaron')
            markets=expected;continue
        if len(cells)==1 and cells[0]==named[1]:year=int(cells[0]);continue
        if cells and cells[0].lower() in r.MONTHS:
            if products is None or markets is None or year is None or len(cells)!=len(products)+1:
                raise r.Error('Columnas internas anuales desconocidas')
            day=date(year,r.MONTHS[cells[0].lower()],1).isoformat()
            if (block,day) in seen or day>as_of.isoformat():raise r.Error('Mes anual duplicado/futuro')
            seen.add((block,day))
            for col,(product,market,raw) in enumerate(zip(products,markets,cells[1:]),1):
                yield record('internal',day,product,market,'Precio interno mensual','ARS','TN',raw,m,
                    f'annual-row:{index};column:{col}',style='ar',flags=['official_annual_monthly_table','currency_context_ARS'],currency_evidence='contextual_annual_internal')
    if blocks!={0,1} or not seen:raise r.Error('Bloques anuales internos incompletos')


def label(value):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD',str(value or ''))
                           if not unicodedata.combining(c)).upper().split())


def numeric(value, style='json'):
    if value is None: return '', 'missing'
    if isinstance(value, bool): raise r.Error('Precio booleano inválido')
    text=str(value).strip()
    if text.upper() in ('S/C','-','--','#N/A',''): return '', 'missing'
    return r.number(text,style)


def record(family, day, product, market, kind, currency, unit, raw, m, locator,
           condition='', style='json', flags=(), **extra):
    value,status=numeric(raw,style)
    row=r.observation(family,day,product,market,kind,currency,unit,value or '',
                      style='json',condition=condition,**extra)
    row.update(price_raw='' if raw is None else str(raw),price=value,value_status=status,
               source_url=m['source_url'],raw_sha256=m['sha256'],capture_id=m['capture_id'],
               updated_at_utc=m['captured_at_utc'],source_locator=locator,
               quality_flags=list(flags),history_schema_version=VERSION)
    row['record_id']=r.identity([row['observation_id'],m['capture_id'],m['sha256'],locator])
    return row


def internal_xlsx(body,m,as_of):
    workbook=load_workbook(io.BytesIO(body),read_only=True,data_only=True)
    if workbook.sheetnames != ['predar t m ysorg','predar s g y lino']:
        raise r.Error('Hojas del histórico interno cambiaron')
    products=[['TRIGO']*5+['MAIZ']*4+['SORGO']*3+['CEBADA FORRAJERA']*2,
              ['SOJA']*4+['GIRASOL']*4+['LINO']*3]
    markets=[['DARSENA','ROSARIO','B.BLANCA','CÓRDOBA','QUEQUEN','DARSENA','ROSARIO','B.BLANCA','QUEQUEN','DARSENA','ROSARIO','B.BLANCA','B.BLANCA','QUEQUEN'],
             ['DARSENA','ROSARIO','B.BLANCA','QUEQUEN','DARSENA','ROSARIO','B.BLANCA','QUEQUEN','DARSENA','ROSARIO','QUEQUEN']]
    for block,sheet in enumerate(workbook):
        title=str(sheet.cell(2,3).value or '')
        if 'EVOLUCION DE PRECIOS INTERNOS' not in title or '$/tn.' not in title:
            raise r.Error('Título/unidad del XLS interno no reconocido')
        heading=[sheet.cell(6,c).value for c in range(3,3+len(products[block]))]
        if [label(v) for v in heading] != [label(v) for v in markets[block]]:
            raise r.Error('Plazas del histórico interno cambiaron')
        year=None; seen=set()
        for index,cells in enumerate(sheet.iter_rows(max_col=16,values_only=True),1):
            key=str(cells[1] or '').strip()
            if re.fullmatch(r'\d{4}',key):year=int(key);continue
            if key.lower() not in r.MONTHS:continue
            if year is None:raise r.Error('Mes interno sin año')
            day=date(year,r.MONTHS[key.lower()],1).isoformat()
            if day>as_of.isoformat():
                if any(v not in (None,'','S/C','s/c') for v in cells[2:2+len(products[block])]):
                    raise r.Error('Precio interno futuro en el XLS')
                continue
            if day in seen:raise r.Error('Mes repetido en hoja interna')
            seen.add(day)
            # Meses futuros en blanco del año vigente no son publicaciones sin precio.
            values=cells[2:2+len(products[block])]
            if all(v is None for v in values):continue
            for col,(product,market,raw) in enumerate(zip(products[block],markets[block],values),3):
                yield record('internal',day,product,label(market),'Precio interno mensual','ARS','TN',raw,m,
                    f'{sheet.title}!R{index}C{col}',flags=['currency_context_ARS','published_monthly'],
                    currency_evidence='contextual_historical_workbook')
    workbook.close()


def board_archive(obj,m,as_of):
    expected=['PRODUCTO','BUENOS AIRES','BAHIA BLANCA','QUEQUEN','ROSARIO','CORDOBA']
    products={'GIRASOL','TRIGO PAN','MAIZ','SOJA','SORGO','CEBADA FORRAJERA'}
    if 'SIN COTIZACION' in label(obj['text']):raise NoPublication('Cámara publica SIN COTIZACIÓN')
    if not obj['text'] or 'No Disponible' in obj['text']:raise Unavailable('Archivo Cámara no disponible')
    if 'Cotizaciones provisorias' not in obj['text']:raise r.Error('Definición de Cámara histórica cambió')
    dates=re.findall(r'(\d{1,2})(?: de)? ([a-z]+)(?: de)? (\d{4})',obj['text'].lower())
    days={date(int(y),r.MONTHS[month],int(d)).isoformat() for d,month,y in dates if month in r.MONTHS}
    named=re.search(r'camaras(\d{4}-\d{2}-\d{2})\.php$',m['source_url'])
    if not named or days!={named[1]} or named[1]>as_of.isoformat():raise r.Error('Fecha Cámara histórica no coincide con enlace')
    active=False;seen=set()
    for index,cells in enumerate(obj['rows']):
        if cells==expected:active=True;continue
        if not active or not cells:continue
        if (not any(cells) or (cells[0]=='Cotizaciones provisorias - Sujetas a ajuste' and not any(cells[1:]))):continue
        if cells[0] in products:
            if len(cells)!=6 or cells[0] in seen:raise r.Error('Columnas/duplicados Cámara histórica')
            seen.add(cells[0])
            for col,(market,raw) in enumerate(zip(expected[1:],cells[1:]),1):
                if re.fullmatch(r'\d+(?:\.\d{3})*,\d+',raw):style='ar'
                elif re.fullmatch(r'\d+(?:,\d{3})*\.\d+',raw):style='us'
                else:style='json'
                yield record('board',named[1],cells[0],market,'Pizarra provisional','ARS','TN',raw,m,
                    f'table-row:{index};column:{col}',style=style,
                    flags=['provisional','currency_unit_context_historical_camera']+(['published_error_marker'] if raw=='#N/A' else []),currency_evidence='contextual')
        elif len(cells)==6:raise r.Error('Producto Cámara desconocido')
    if seen!=products:raise r.Error('Bloques Cámara incompletos')


def fas_xlsx(body,m,as_of):
    w=load_workbook(io.BytesIO(body),read_only=True,data_only=True)
    if len(w.sheetnames)!=1:raise r.Error('Hojas export FAS inesperadas')
    s=w.active;regime=m['regime']
    title='FAS - Derechos de Exportacion '+('Completos' if regime=='D.E.C.' else 'Reducidos')
    if label(s.cell(1,1).value)!=label(title):raise r.Error('Régimen FAS no coincide con export')
    expected=['Fecha']+r.FAS_LABELS
    headers=[i for i in (6,7) if [label(s.cell(i,c).value) for c in range(1,11)] == [label(v) for v in expected]]
    if len(headers)!=1:
        raise r.Error('Columnas export FAS cambiaron')
    seen=set()
    for index,cells in enumerate(s.iter_rows(min_row=headers[0]+1,max_col=10,values_only=True),headers[0]+1):
        if not any(v is not None for v in cells):continue
        if not isinstance(cells[0],datetime):raise r.Error('Fecha export FAS inválida')
        day=cells[0].date().isoformat()
        if day in seen or not m['requested_from']<=day<=min(m['requested_to'],as_of.isoformat()):raise r.Error('Fecha duplicada/fuera de rango FAS')
        seen.add(day)
        for col,(code,product,raw) in enumerate(zip(r.FAS_CODES,r.FAS_LABELS,cells[1:]),2):
            known=code not in ('AS','AG')
            flags=['official_export', 'unit_contextual_grain' if known else 'currency_unit_unknown']
            if code=='GI':flags.append('sunflower_formula_pre_2019_02_14' if day<'2019-02-14' else 'sunflower_formula_from_2019_02_14')
            if day>='2025-04-14':flags.append('published_exchange_rate_correction_note')
            yield record('fas',day,product,'Argentina / paridad teórica',f'FAS teórico {regime}',
                'ARS' if known else 'Sin identificar','TN' if known else 'Sin identificar',raw,m,
                f'{s.title}!R{index}C{col}',condition=regime,flags=flags,regime=regime,product_code=code,
                currency_evidence='contextual' if known else 'unknown')
    w.close()


def futures_matrix(obj,m,as_of):
    # El título dice promedios mensuales, pero las columnas identifican días reales
    # y las filas posiciones. Se conservan ambos; la columna Promedio no se trata como día.
    text=label(obj['text'])
    match=re.search(r'MERCADO (?:DE )?(CHICAGO|KANSAS) - PROMEDIOS MENSUALES EN DOL/TON\s+(.+)',text)
    if not match:
        # Variante publicada de 1997: el encabezado de página identifica el mercado.
        match=re.search(r'MERCADO DE (CHICAGO|KANSAS) - .+? MERCADO - PROMEDIOS MENSUALES EN DOL/TON\s+(.+)',text)
    if not match:raise r.Error('Definición de matriz futuros no reconocida')
    market=match[1]
    known={'maiz':'MAIZ','har_soja':'HARINA DE SOJA','ac_soja':'ACEITE DE SOJA',
           'soja':'SOJA','avena':'AVENA','trigo_pan':'TRIGO PAN'}
    filename=m['source_url'].rsplit('/',1)[-1]
    named=re.fullmatch(r'\d+_([a-z_]+)_(\d{4})\.php',filename)
    if not named or named[1] not in known:raise r.Error('Producto/año matriz no reconocido')
    product=known[named[1]]
    if not match[2].startswith(product+' '):raise r.Error('Producto del título no coincide con archivo')
    dates=None;seen={}
    for index,cells in enumerate(obj['rows']):
        if cells and label(cells[0])=='INICIO EMBARQ':
            dates=[];summary=False
            for col,cell in enumerate(cells[1:],1):
                if re.fullmatch(r'\d{2}/\d{2}/(?:\d{4}|\d{2})',cell):
                    parts=cell.split('/')
                    if len(parts[2])==2:
                        if parts[2]!=named[2][-2:]:raise r.Error('Año abreviado no coincide con archivo')
                        parts[2]=named[2]
                    day=date(int(parts[2]),int(parts[1]),int(parts[0])).isoformat()
                    if day[:4]!=named[2] or day>as_of.isoformat():raise r.Error('Fecha matriz fuera del año')
                    dates.append((col,day))
                elif label(cell)=='PROMEDIO':summary=True
                elif label(cell)=='' or (summary and cell=='-'):continue
                else:raise r.Error('Encabezado matriz desconocido')
            continue
        if cells and re.fullmatch(r'[A-Z]{3}\d{4}',cells[0]):
            if not dates:raise r.Error('Posición sin fechas en matriz')
            position=cells[0];month=r.SHORT.get(position[:3].lower())
            if not month:raise r.Error('Vencimiento matriz no reconocido')
            maturity=date(int(position[3:]),month,1).isoformat()[:7]
            for col,day in dates:
                if col>=len(cells):raise r.Error('Fila matriz truncada')
                key=(day,position)
                raw=cells[col]
                if key in seen:
                    if seen[key]==raw:continue
                    raise r.Error('Cotización matriz duplicada con valores diferentes')
                seen[key]=raw
                if re.fullmatch(r'\d+\.\d{3}',raw):raise r.Error('Separador decimal ambiguo en matriz')
                # Algunas matrices mezclan coma y punto decimal. Un punto seguido de
                # tres cifras es ambiguo: no se interpreta sin evidencia adicional.
                style='json' if re.fullmatch(r'\d+\.\d{1,2}',raw) else 'ar'
                yield record('futures',day,product,market,'Futuro / histórico publicado','USD','TN',raw,m,
                    f'table-row:{index};column:{col}',condition=position,style=style,flags=['historical_matrix_quote','not_a_continuous_contract'],
                    position_raw=position,shipment_window=maturity,currency_evidence='documented')


def futures_daily(obj,m,as_of):
    named=re.search(r'cotiza_(\d{4}-\d{2}-\d{2})\.php$',m['source_url'])
    if not named or r.spanish_date(obj['text'])!=named[1] or named[1]>as_of.isoformat():
        raise r.Error('Fecha del informe diario futuros no coincide con archivo')
    # Variante literal verificada de los informes antiguos. Cambia únicamente
    # mayúsculas/acentos del contexto y la abreviatura Tn.→Ton.; no toca precios,
    # fechas, posiciones, encabezados de productos ni el parser corriente.
    old='* En Dolares Estadounidenses/Tn.'
    variant='COTIZACIONES DE CIERRE' in obj['text'] and old in obj['text']
    prepared={**obj,'text':obj['text'].replace('COTIZACIONES DE CIERRE','Cotizaciones de cierre').replace(old,'* En Dólares Estadounidenses/Ton.')} if variant else obj
    rows=r.parse_futures(prepared);seen=set()
    if not rows:raise r.Error('Informe de futuros vacío')
    for row in rows:
        if row['observation_id'] in seen:raise r.Error('Informe de futuros duplicado')
        seen.add(row['observation_id'])
        row.update(source_url=m['source_url'],raw_sha256=m['sha256'],capture_id=m['capture_id'],
            updated_at_utc=m['captured_at_utc'],source_locator='daily-report:'+row['observation_id'],
            quality_flags=['official_daily_closing_report','not_a_continuous_contract']+(['historical_context_spelling_variant'] if variant else []),history_schema_version=VERSION)
        row['record_id']=r.identity([row['observation_id'],m['capture_id'],m['sha256']])
        yield row


def fob_sspm(body,m,distribution,as_of):
    fields=distribution['field'];dictionary={x['title']:x for x in fields}
    reader=csv.DictReader(io.StringIO(body.decode('utf-8-sig')))
    if set(reader.fieldnames or [])!=set(dictionary) or reader.fieldnames[0]!='indice_tiempo':
        raise r.Error('Schema FOB SSPM distinto del catálogo oficial')
    if any(x.get('units')!='Dolares por tonelada' for k,x in dictionary.items() if k!='indice_tiempo'):
        raise r.Error('Unidad FOB SSPM desconocida')
    seen=set()
    for index,cells in enumerate(reader,2):
        day=date.fromisoformat(cells['indice_tiempo'])
        if day>as_of or day in seen:raise r.Error('Índice temporal FOB duplicado/futuro')
        seen.add(day)
        for column,raw in cells.items():
            if column=='indice_tiempo':continue
            field=dictionary[column]
            flags=['official_sspm_time_series','index_date_not_confirmed_publication_date',
                   'no_mapping_to_api_position_or_circular','no_shipment_window']
            if day.weekday()>=5:flags.append('weekend_value_published_by_sspm')
            yield record('fob',day.isoformat(),field['description'],'FOB oficial / serie SSPM',
                'FOB oficial / serie SSPM','USD','TN',raw,m,f'CSV-row:{index};field:{column}',
                condition='Serie SSPM; embarque y circular no informados',flags=flags,
                product_code=field['id'],currency_evidence='documented_sspm_catalog')
