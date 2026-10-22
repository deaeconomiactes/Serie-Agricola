"""Read linked MAGyP cattle workbooks without evaluating or manufacturing values.

Each worksheet/block remains an independent published table. Monthly cells and
published cumulative/share rollups coexist with explicit grain and flags; they
must never be summed together. Classification labels on either side of 2019 are
not bridged. Formula cells without a cached value remain missing.
"""
from __future__ import annotations

import calendar
import io
import re
import unicodedata
from datetime import date
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from scripts.carnes.model import DataError, observation

MONTHS = {name: i for i, name in enumerate(
    ['ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO',
     'SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE'], 1)}
CATEGORIES = {'BUEY':'Buey','NOVILLITO':'Novillito','NOVILLO':'Novillo',
    'TERNERA':'Ternera','TERNERO':'Ternero','TORITO':'Torito','TORO':'Toro',
    'VAQUILLONA':'Vaquillona','VACA':'Vaca','M.E.J.':'M.E.J.',
    'TOTAL GENERAL':'Total general','TOTAL':'Total general'}
AGES = {'2 DIENTES','4 DIENTES','6 DIENTES','8 DIENTES','MAS DE 8 DIENTES'}
COVERS = {'PRESENTACION','CARATULA'}
OLD_CATEGORIES={'Buey','Novillito','Novillo','Ternera','Ternero','Torito','Toro','Vaquillona','Vaca'}
MODERN_CELLS={('M.E.J.','2 DIENTES'),('Novillito','2 DIENTES'),('Novillito','4 DIENTES'),
    ('Novillo','6 DIENTES'),('Novillo','8 DIENTES'),('Novillo','MAS DE 8 DIENTES'),
    ('Toro','4 DIENTES'),('Toro','6 DIENTES'),('Toro','8 DIENTES'),('Toro','MAS DE 8 DIENTES'),
    ('Vaquillona','2 DIENTES'),('Vaquillona','4 DIENTES'),
    ('Vaca','6 DIENTES'),('Vaca','8 DIENTES'),('Vaca','MAS DE 8 DIENTES')}


def label(value):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD',str(value or ''))
        if not unicodedata.combining(c)).upper().split())


def raw_text(value):
    return '' if value is None else str(value)


def numeric(value, original):
    flags=[]
    if value is None or value == '':
        flags.append('formula_result_not_cached' if isinstance(original,str) and original.startswith('=') else 'missing_published_value')
        return None,flags
    if isinstance(value,bool):raise DataError('Boolean is not a published quantity')
    if isinstance(value,str) and value.strip() in {'-','--'}:
        return None,['not_reported_marker']
    try:
        result=Decimal(str(value))
    except (InvalidOperation,ValueError,TypeError):
        raise DataError('Non-numeric published quantity') from None
    if not result.is_finite() or result<0:raise DataError('Invalid published quantity')
    if result != result.to_integral_value():flags.append('fractional_published_value')
    return format(result,'f'),flags


def identifier(value):
    if value is None:return ''
    if isinstance(value,bool):raise DataError('Invalid entity identifier')
    if isinstance(value,(float,int)):
        number=Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value():raise DataError('Fractional entity identifier')
        return str(int(number))
    return str(value).strip()


def merged_value(sheet,row,column):
    cell=sheet.cell(row,column)
    if cell.value is not None:return cell.value
    for area in sheet.merged_cells.ranges:
        if area.min_row<=row<=area.max_row and area.min_col<=column<=area.max_col:
            return sheet.cell(area.min_row,area.min_col).value
    return None


def title_year(text,as_of,expected_year):
    match=re.match(r'^ANO\s+(\d{4})\b',label(text))
    if not match:raise DataError('Missing workbook observation year')
    year=int(match.group(1))
    if not 1900<=year<=as_of.year or (expected_year is not None and year!=expected_year):
        raise DataError('Workbook year does not match requested period')
    return year


def period(year,month,as_of):
    start=date(year,month,1)
    if start>as_of:raise DataError('Future workbook observation')
    return start.isoformat(),date(year,month,calendar.monthrange(year,month)[1]).isoformat()


def add_cell(out,sheet,formulas,row,column,*,year,month,as_of,flags=None,**dimensions):
    value=sheet.cell(row,column).value
    original=formulas.cell(row,column).value
    parsed,value_flags=numeric(value,original)
    start,end=period(year,month,as_of)
    if dimensions.get('frequency') in {'year_to_date','annual'}:start=f'{year}-01-01'
    unit=dimensions.get('unit','heads')
    if unit=='share' and parsed is not None and Decimal(parsed)>1:
        raise DataError('Published participation is outside [0,1]')
    quality=['published_table',*(flags or []),*value_flags]
    if end>as_of.isoformat():quality.append('period_in_progress')
    frequency=dimensions.pop('frequency','monthly')
    out.append(observation(observation_date=start,period_end=end,frequency=frequency,
        source_locator=f'{sheet.title}!{sheet.cell(row,column).coordinate}',
        value=parsed,value_raw=raw_text(original),quality_flags=quality,**dimensions))


def schema_cells(sheet,header,known_columns,start_row,end_row,max_column=None):
    """Unexpected data columns are schema errors, not silently dropped dimensions."""
    last=max(known_columns)
    end=sheet.max_column if max_column is None else max_column
    if last>=end:return
    for row in sheet.iter_rows(min_row=start_row,max_row=end_row,min_col=last+1,max_col=end):
        if any(cell.value is not None for cell in row):
            raise DataError(f'Unknown populated column in {sheet.title}')


def table_header(sheet):
    for row in sheet.iter_rows(min_row=1,max_row=min(8,sheet.max_row)):
        months={cell.column:MONTHS[label(cell.value)] for cell in row if label(cell.value) in MONTHS}
        if months:
            if list(months.values()) != list(range(1,max(months.values())+1)):
                raise DataError('Unexpected month columns')
            return row[0].row,months
    raise DataError('Missing monthly table header')


def grid_table(sheet,formulas,kind,as_of,expected_year,species):
    title=sheet.cell(1,1).value
    year=title_year(title,as_of,expected_year)
    if 'FAENA DE BOVINOS' not in label(title) or 'EN CABEZAS' not in label(title):
        raise DataError('Unknown slaughter table context')
    header,months=table_header(sheet)
    first=min(months);month_end=max(months.values())
    expected_first={'province':2,'establishment':5,'slaughter_user':4}[kind]
    if first!=expected_first:raise DataError('Entity table columns changed')
    names=[label(sheet.cell(header,column).value) for column in range(1,first)]
    if kind=='province' and names not in [['PROVINCIA'],['PROVINCIA/MES']]:raise DataError('Province header changed')
    if kind=='establishment' and not (names[0] in {'N° EST.','No EST.','NO EST.'} and names[1].startswith('CUIT') and names[2]=='RAZON SOCIAL FRIGORIFICO' and names[3]=='PROVINCIA'):
        raise DataError('Establishment header changed')
    if kind=='slaughter_user' and not (names[0].startswith('CUIT') and names[1] in {'RAZON SOCIAL USUARIO','RAZON SOCIAL USUARIO DE FAENA'} and names[2]=='ACTIVIDAD'):
        raise DataError('Slaughter user header changed')
    extras={};quarantine_flags=[]
    for cell in sheet[header]:
        text=label(cell.value)
        if cell.column<=max(months):continue
        if not text:continue
        if text==f'ACUMULADO {year}':extras[cell.column]='cumulative'
        elif text==f'TOTAL ANO {year}' and month_end==12:extras[cell.column]='annual'
        elif (year==2019 and kind=='establishment' and cell.coordinate=='Q3' and text=='ACUMULADO 2020'):
            # Verified defect in the linked 2019 archive. Its dated monthly cells
            # are usable; this conflicting aggregate is retained only in RAW.
            extras[cell.column]='quarantined_cumulative'
            quarantine_flags.append(f'quarantined_aggregate_header:{sheet.title}!{cell.coordinate}:{raw_text(cell.value)}')
        elif kind=='province' and text.startswith('PARTICIPACION') and 'FAENA' in text:extras[cell.column]='share'
        else:raise DataError('Unknown aggregate column')
    if sum(value in {'cumulative','annual','quarantined_cumulative'} for value in extras.values())!=1:raise DataError('Missing cumulative column')
    out=[];data_rows=0;max_column=sheet.max_column
    for row in range(header+1,sheet.max_row+1):
        metadata=[sheet.cell(row,column).value for column in range(1,first)]
        texts=[label(value) for value in metadata]
        if any(text.startswith('FUENTE:') for text in texts):break
        numbers=[sheet.cell(row,column).value for column in [*months,*extras]]
        if not any(metadata) and not any(value is not None for value in numbers):continue
        # Older files have a blank merged second header row.
        if not any(metadata) and row==header+1 and not any(value is not None for value in numbers):continue
        total=any(text in {'TOTAL','TOTALES','TOTAL GENERAL','TOTAL PAIS'} for text in texts)
        unlabelled=not any(metadata)
        if unlabelled:total=True
        if total:
            dimensions=dict(entity_type='total' if not unlabelled else 'unlabelled_published_row',
                entity_id='',entity_name=next((raw_text(v) for v in metadata if v is not None),'Unlabelled published row'),
                geography='Argentina' if not unlabelled else 'NOT_IDENTIFIED',geography_level='national' if not unlabelled else 'not_identified')
        elif kind=='province':
            dimensions=dict(entity_type='province',entity_id='',entity_name='',
                geography=raw_text(metadata[0]).strip(),geography_raw=raw_text(metadata[0]),geography_level='province')
        elif kind=='establishment':
            if not all(metadata):raise DataError('Incomplete establishment identity')
            dimensions=dict(entity_type='establishment',entity_id=identifier(metadata[0]),
                entity_tax_id=identifier(metadata[1]),entity_name=raw_text(metadata[2]),
                geography=raw_text(metadata[3]).strip(),geography_raw=raw_text(metadata[3]),geography_level='province')
        else:
            if not all(metadata):raise DataError('Incomplete slaughter-user identity')
            dimensions=dict(entity_type='slaughter_user',entity_id=identifier(metadata[0]),
                entity_tax_id=identifier(metadata[0]),entity_name=raw_text(metadata[1]),entity_activity=raw_text(metadata[2]),
                geography='Argentina',geography_level='national')
        data_rows+=int(not total)
        base=dict(family='slaughter',metric='slaughter_heads',species=species,
            category='All categories',category_raw='Todas las categorías',unit='heads',unit_raw='CABEZAS',
            weight_basis='NOT_APPLICABLE',
            methodology='MAGyP published entity aggregate; categories not reconstructed',
            table_id=f'slaughter/{kind}',**dimensions)
        flags=[*quarantine_flags,*(['published_rollup'] if total else [])]
        if unlabelled:flags.append('unlabelled_numeric_row')
        if species=='bovinos_y_bubalinos':flags.append('combined_species_not_separable')
        for column,month in months.items():
            add_cell(out,sheet,formulas,row,column,year=year,month=month,as_of=as_of,flags=flags,**base)
        for column,aggregate in extras.items():
            if aggregate=='quarantined_cumulative':continue
            extra={**base,'frequency':'annual' if aggregate=='annual' else 'year_to_date'}
            if aggregate=='share':extra.update(metric='slaughter_share',unit='share',unit_raw=raw_text(sheet.cell(header,column).value))
            add_cell(out,sheet,formulas,row,column,year=year,month=month_end,as_of=as_of,
                flags=[*flags,'published_rollup',{'cumulative':'published_cumulative','annual':'published_annual','share':'published_share'}[aggregate]],**extra)
        schema_cells(sheet,header,[*months,*extras],row,row,max_column=max_column)
    if not data_rows:raise DataError('Empty slaughter detail table')
    if year==2020 and kind=='establishment':
        # The official archive repeats these exact identities without a dimension
        # separating the rows. One pair has conflicting October values, both have
        # different cumulative totals. Do not sum, select a winner or invent IDs.
        # Quarantine the complete affected entity groups; RAW retains every cell.
        defects=[('422','30716396696','CARNES ESPECIALES ENTRERRIANAS SRL','ENTRE RIOS',{216,270}),
                 ('47','30607813295','LE GURIE S.A.','BUENOS AIRES',{219,226})]
        excluded=set();scope=[]
        for entity,tax,name,province,expected_rows in defects:
            matching=[item for item in out if item['entity_id']==entity and item['entity_tax_id']==tax
                and item['entity_name']==name and item['geography']==province]
            source_rows={int(re.search(r'(\d+)$',item['source_locator']).group(1)) for item in matching}
            if source_rows==expected_rows:
                excluded.update(item['source_locator'] for item in matching)
                locators=','.join(f'A{number}:Q{number}' for number in sorted(expected_rows))
                scope.append(f'quarantined_conflicting_entity_rows:{sheet.title}!{locators}:entity_id={entity}')
        if excluded:
            out=[item for item in out if item['source_locator'] not in excluded]
            for item in out:item['quality_flags'].extend(scope)
    return out


def category_blocks(sheet,formulas,as_of,expected_year,species,production=False):
    headers=[row[0].row for row in sheet.iter_rows() if label(row[0].value).startswith('MES /') or label(row[0].value).startswith('MES/')]
    if not headers:raise DataError('Missing category header')
    if not production and len(headers)!=1:raise DataError('Unexpected category blocks')
    out=[];families=set()
    for block,header in enumerate(headers):
        next_header=headers[block+1] if block+1<len(headers) else sheet.max_row+1
        titles=[sheet.cell(row,1).value for row in range(max(1,header-3),header) if label(sheet.cell(row,1).value).startswith('ANO ')]
        if not titles:raise DataError('Missing category title')
        title=titles[-1];context=label(title);year=title_year(title,as_of,expected_year)
        if 'EN TONELADAS RES CON HUESO' in context:
            family,metric,unit,unit_raw='production','carcass_production','tonnes_carcass_bone_in','TONELADAS RES CON HUESO'
        elif 'EN KG RES CON HUESO' in context:
            family,metric,unit,unit_raw='carcass_weight','mean_carcass_weight','kg_carcass_bone_in','KG RES CON HUESO'
        elif 'EN CABEZAS' in context:
            family,metric,unit,unit_raw='slaughter','slaughter_heads','heads','CABEZAS'
        else:raise DataError('Unknown category measure context')
        if not production and family!='slaughter':raise DataError('Wrong slaughter family')
        if family in families:raise DataError('Repeated measure block')
        families.add(family)
        month_rows=[row for row in range(header+1,next_header) if label(sheet.cell(row,1).value) in MONTHS]
        if not month_rows:raise DataError('Empty category block')
        months=[MONTHS[label(sheet.cell(row,1).value)] for row in month_rows]
        if months!=list(range(1,max(months)+1)):raise DataError('Repeated or missing category month')
        columns={}
        for column in range(2,sheet.max_column+1):
            raw=merged_value(sheet,header,column)
            if raw is None:continue
            name=label(raw)
            if name not in CATEGORIES:raise DataError('Unknown published cattle category')
            age_raw=sheet.cell(header+1,column).value if header+1 not in month_rows else None
            age=label(age_raw)
            if age and age not in AGES:raise DataError('Unknown published dentition')
            if name.startswith('TOTAL'):method='published_all_classifications_total'
            elif age:method='resolution_32_2018_dentition'
            else:method='pre_2019_weight_sex_categories'
            if year>2019 and not age and not name.startswith('TOTAL'):raise DataError('Missing dentition in modern category')
            columns[column]=(CATEGORIES[name],raw_text(raw),raw_text(age_raw),method)
        if not columns or sum(name=='Total general' for name,_,_,_ in columns.values())!=1:
            raise DataError('Missing category total or category columns')
        old=[category for category,_,age,_ in columns.values() if not age and category!='Total general']
        modern=[(category,label(age)) for category,_,age,_ in columns.values() if age]
        if ((year<=2019 and (len(old)!=9 or set(old)!=OLD_CATEGORIES)) or
                (year>2019 and old) or
                (year>=2019 and (len(modern)!=15 or set(modern)!=MODERN_CELLS)) or
                (year<2019 and modern)):
            raise DataError('Incomplete category/dentition schema')
        schema_cells(sheet,header,columns,min(month_rows),max(month_rows))
        base=dict(family=family,metric=metric,species=species,unit=unit,unit_raw=unit_raw,
            geography='Argentina',geography_level='national',entity_type='category',
            weight_basis='carcass_bone_in' if family!='slaughter' else 'NOT_APPLICABLE',
            table_id=f'{family}/category_age')
        for row in month_rows:
            for column,(category,category_raw,age,method) in columns.items():
                flags=[]
                if year==2019:flags.append('classification_transition_2019_no_bridge')
                if species=='bovinos_y_bubalinos':flags.append('combined_species_not_separable')
                if category=='Total general':flags.append('published_rollup')
                add_cell(out,sheet,formulas,row,column,year=year,month=MONTHS[label(sheet.cell(row,1).value)],
                    as_of=as_of,flags=flags,category=category,category_raw=category_raw,age_class=age,
                    methodology=method,**base)
        # Preserve published cumulative/category/share rows, not formulas-derived totals.
        for row in range(max(month_rows)+1,next_header):
            text=label(sheet.cell(row,1).value)
            if not (text.startswith('TOTAL POR ') or text.startswith('PARTICIPACION EN EL TOTAL') or text.startswith('PESO MEDIO DE LA RES POR ')):
                continue
            category_rollup='CATEGORIA' in text or text.startswith('PARTICIPACION')
            flags=['published_rollup','published_share' if text.startswith('PARTICIPACION') else 'published_cumulative']
            for column,(category,category_raw,age,method) in columns.items():
                # A merged category total occupies only its first cell; never replicate it.
                if category_rollup:
                    if isinstance(sheet.cell(row,column),MergedCell):continue
                    if formulas.cell(row,column).value is None:continue
                    age=''
                extra={**base,'frequency':'year_to_date','entity_type':'category_rollup',
                    'table_id':f'{family}/category_age/{text}',
                    'category':category,'category_raw':category_raw,'age_class':age,'methodology':method}
                if text.startswith('PARTICIPACION'):
                    extra.update(metric=f'{family}_share',unit='share',unit_raw=raw_text(sheet.cell(row,1).value))
                add_cell(out,sheet,formulas,row,column,year=year,month=max(months),as_of=as_of,flags=flags,**extra)
    if production and families!={'slaughter','production','carcass_weight'}:
        raise DataError('Incomplete production/weight workbook')
    return out


def parse_workbook(payload:bytes,kind:str,as_of:date,expected_year:int|None=None)->list[dict]:
    if kind not in {'slaughter_detail','production_detail'}:raise DataError('Unknown workbook parser kind')
    if not payload.startswith(b'PK') or len(payload)>10_000_000:raise DataError('Invalid XLSX payload')
    try:
        book=load_workbook(io.BytesIO(payload),data_only=True)
        formulas=load_workbook(io.BytesIO(payload),data_only=False)
    except Exception as error:
        raise DataError('Unreadable XLSX payload') from error
    try:
        if any(sheet.max_row>10000 or sheet.max_column>300 for sheet in book):raise DataError('Workbook outside size contract')
        covers=' '.join(label(cell.value) for sheet in book if label(sheet.title) in COVERS for row in sheet for cell in row if cell.value is not None)
        species='bovinos_y_bubalinos' if 'BOVINOS Y BUBALINOS' in covers else 'bovinos'
        out=[];tables=set()
        for sheet in book:
            name=label(sheet.title)
            if name in COVERS:continue
            if kind=='production_detail':
                if name not in {'PRODUCCION POR CATEGORIA Y EDAD','FAENA Y PRODUCCION POR CATEGORIA Y EDAD'}:raise DataError('Unknown production worksheet')
                out.extend(category_blocks(sheet,formulas[sheet.title],as_of,expected_year,species,production=True));continue
            kinds={'POR PROVINCIA':'province','POR ESTABLECIMIENTO':'establishment','POR USUARIO DE FAENA':'slaughter_user'}
            if name in kinds:
                tables.add(kinds[name]);out.extend(grid_table(sheet,formulas[sheet.title],kinds[name],as_of,expected_year,species))
            elif name in {'POR CATEGORIA','POR CATEGORIA Y EDAD'}:
                tables.add('category');out.extend(category_blocks(sheet,formulas[sheet.title],as_of,expected_year,species))
            else:raise DataError('Unknown slaughter worksheet')
        if kind=='slaughter_detail' and tables!={'province','establishment','slaughter_user','category'}:
            raise DataError('Incomplete slaughter workbook')
        if not out or not any(row['value_status']=='observed' for row in out):raise DataError('Workbook has no observed quantities')
        shared_flags={flag for row in out for flag in row['quality_flags'] if flag.startswith(
            ('quarantined_aggregate_header:','quarantined_conflicting_entity_rows:'))}
        if any('PROVISOR' in label(cell.value) for sheet in book for row in sheet for cell in row if isinstance(cell.value,str)):
            shared_flags.add('provisional_source_statement')
        if shared_flags:
            for row in out:row['quality_flags']=list(dict.fromkeys([*row['quality_flags'],*sorted(shared_flags)]))
        return out
    finally:
        book.close();formulas.close()
