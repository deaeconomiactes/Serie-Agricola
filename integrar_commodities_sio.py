#!/usr/bin/env python3
"""Integra respuestas reales de SIO Granos, separadas del pipeline BCR."""

from __future__ import annotations

import csv
import json
import re
import unicodedata
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable


ROOT = Path(__file__).resolve().parent
RAW_DIR = ROOT / "data" / "commodities_sio" / "raw"
PROCESSED_DIR = ROOT / "data" / "commodities_sio" / "processed"
DASHBOARD_DIR = ROOT / "data" / "commodities_sio" / "dashboard"
CATALOG_PATH = ROOT / "data" / "commodities_sio" / "catalogo_productos_sio.csv"
MAPPING_LOCAL_PATH = ROOT / "data" / "commodities_sio" / "mapeo_getoperaciones_sio.local.json"
MAPPING_EXAMPLE_PATH = ROOT / "data" / "commodities_sio" / "mapeo_getoperaciones_sio.example.json"
OUTPUT_PATH = PROCESSED_DIR / "COMMODITIES_SIO_INTEGRADO.csv"
PAGINATED_OUTPUT_PATH = PROCESSED_DIR / "COMMODITIES_SIO_MUESTRA_PAGINADA.csv"
MANUAL_EXPORT_OUTPUT_PATH = PROCESSED_DIR / "COMMODITIES_SIO_EXPORTACION_MANUAL.csv"
LATEST_SNAPSHOT_OUTPUT_PATH = PROCESSED_DIR / "COMMODITIES_SIO_LATEST_SNAPSHOT.csv"
HISTORIC_SNAPSHOTS_OUTPUT_PATH = PROCESSED_DIR / "COMMODITIES_SIO_HISTORICO_SNAPSHOTS.csv"
LIGHT_HISTORY_OUTPUT_PATH = DASHBOARD_DIR / "COMMODITIES_SIO_HISTORICO_SNAPSHOTS_LIVIANO.csv"
PAGINATED_REPORT_PATH = ROOT / "data" / "commodities_sio" / "reports" / "REPORTE_MUESTRA_PAGINADA_SIO.md"
PAGINATION_REPORT_PATH = ROOT / "data" / "commodities_sio" / "reports" / "REPORTE_PAGINACION_SIO.md"
OUTPUT_COLUMNS = [
    "fecha", "año", "mes", "commodity", "fuente", "mercado", "tipo_precio",
    "precio_tipo_original", "precio_unidad", "precio_total", "campo_precio_original", "valor_precio_original", "precio_original_texto", "moneda", "moneda_explicitamente_informada", "moneda_inferida", "campo_moneda_original", "valor_moneda_original", "unidad", "precio", "volumen", "volumen_unidad", "campo_volumen_original", "procedencia",
    "provincia", "localidad", "zona", "lugar_entrega", "precio_puesto_en", "operacion",
    "condicion_pago", "condicion_comercial", "frecuencia", "archivo_origen",
    "fecha_integracion", "observaciones", "apto_piloto", "apto_dashboard", "pagina_origen", "id_operacion_sio", "muestra_tipo", "muestra_paginas", "estado_paginacion",
    "precio_cero_flag", "precio_cero_tipo", "precio_valido_para_serie",
]
SNAPSHOT_COLUMNS = ["fecha_descarga_snapshot"] + OUTPUT_COLUMNS
LIGHT_HISTORY_COLUMNS = [
    "id_operacion_sio", "fecha", "año", "mes", "commodity", "fuente", "mercado", "tipo_precio", "operacion",
    "moneda", "unidad", "precio", "precio_original_texto", "volumen", "volumen_unidad", "procedencia",
    "lugar_entrega", "condicion_comercial", "precio_cero_flag", "precio_valido_para_serie", "fecha_descarga_snapshot",
    "fecha_actualizacion_dashboard", "observaciones",
]
LATEST_SNAPSHOT_PATTERN = re.compile(r"^SIO_latest_GetOperaciones_\d{8}_\d{6}\.json$", flags=re.I)
EXTENSIONS = {".json", ".csv", ".xlsx", ".xls", ".html", ".htm"}
NON_REAL_MARKERS = ("plantilla", "simul", "prueba", "ejemplo", "sample")
DEFAULT_SOURCE = "SIO Granos / Secretaría de Agricultura"

# Mapeo respaldado por la estructura de la grilla documentada en
# reports/REPORTE_MAPEO_GETOPERACIONES_SIO.md. Se aplica sólo al endpoint
# GetOperaciones usado por --update-latest; no habilita paginación.
LATEST_POSITIONAL_MAPPING = {
    0: {"source_label": "ID", "target_field": "id_operacion_sio"},
    1: {"source_label": "Fecha Concertación", "target_field": "fecha"},
    3: {"source_label": "Operación", "target_field": "operacion"},
    4: {"source_label": "Tipo", "target_field": "tipo_operacion"},
    5: {"source_label": "Precio", "target_field": "tipo_precio"},
    6: {"source_label": "Producto", "target_field": "commodity"},
    7: {"source_label": "Cant. (TN)", "target_field": "volumen", "unit_field": "volumen_unidad", "unit_value_if_label_matches": "TN"},
    9: {"source_label": "Procedencia Pcia./LOCALID.", "target_field": "procedencia"},
    10: {"source_label": "Precio/TN Monto", "target_field": "precio", "unit_field": "unidad", "unit_value_if_label_matches": "TN"},
    11: {"source_label": "Lugar Entrega", "target_field": "lugar_entrega"},
    13: {"source_label": "Condición Pago", "target_field": "condicion_pago"},
}
HEADER_KEYS = {
    "fecha", "fechadeclaracion", "fechaconcertacion", "fechadeentrega", "producto",
    "grano", "commodity", "especie", "precio", "preciomonto", "preciotn", "monto",
    "cotizacion", "valor", "moneda", "unidad", "volumen", "cantidad", "cantidadtn",
    "tn", "procedencia", "provincia", "localidad", "zona", "operacion",
}


def key(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", normalized.lower())


def text(value: Any) -> str:
    if value is None:
        return ""
    result = str(value).replace("\r", "").strip()
    return "" if result.lower() in {"nan", "nat", "none"} else result


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = text(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            continue
    return None


def parse_number(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    raw = text(value).replace("$", "").replace("ARS", "").replace("USD", "")
    raw = re.sub(r"[^0-9,.-]", "", raw)
    if not raw or raw in {"-", ".", ","}:
        return None
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(",", "") if raw.count(",") > 1 else raw.replace(",", ".")
    elif "." in raw:
        parts = raw.split(".")
        if len(parts) > 1 and all(len(part) == 3 for part in parts[1:]):
            raw = "".join(parts)
    try:
        return float(raw)
    except ValueError:
        return None


def extract_explicit_currency(value: Any) -> tuple[str, str, str]:
    """Normaliza sólo símbolos/textos monetarios presentes en el valor original."""

    raw = text(value)
    if re.search(r"U\s*\$\s*S|US\s*\$|\bUSD\b", raw, flags=re.I):
        return "USD", "sí", "no"
    if "$" in raw:
        return "ARS", "sí", "no"
    return "Sin especificar", "no", "no"


def normalize_explicit_currency_field(value: Any) -> str:
    raw = text(value)
    if not raw:
        return ""
    if raw.strip() in {"0", "-", "N/A", "NA"}:
        return ""
    currency, explicit, _ = extract_explicit_currency(raw)
    if explicit == "sí":
        return currency
    if re.fullmatch(r"USD|US DOLLARS?|DOLARES?|DÓLARES?", raw, flags=re.I):
        return "USD"
    if re.fullmatch(r"ARS|PESOS?|PESOS? ARGENTINOS?", raw, flags=re.I):
        return "ARS"
    return raw


def classify_zero_price(raw_value: Any, parsed_price: float | None, operation: Any = "") -> str:
    """Clasifica el origen observable de un cero sin convertirlo ni eliminarlo."""

    if parsed_price != 0:
        return "no_aplica"
    raw = text(raw_value)
    normalized = key(raw)
    operation_key = key(operation)
    if not raw:
        return "campo_vacio_parseado_cero"
    if normalized in {"sc", "sincotizacion", "sinprecio", "afijar", "fijar"} or any(token in normalized for token in ("sincot", "sinprecio", "afijar")):
        return "texto_sin_precio"
    if any(token in operation_key for token in ("anulacion", "rectificacion")) and normalized in {"0", "00", "000"}:
        return "operacion_sin_precio"
    stripped = re.sub(r"[^0-9]", "", raw)
    if stripped and set(stripped) == {"0"}:
        return "cero_explicito"
    if raw:
        return "parsing_fallido"
    return "no_determinado"


def read_aliases() -> dict[str, str]:
    aliases: dict[str, str] = {}
    if not CATALOG_PATH.exists():
        return aliases
    with CATALOG_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle, delimiter=";"):
            canonical = text(row.get("commodity"))
            for alias in [canonical] + text(row.get("aliases")).split("|"):
                if alias:
                    aliases[key(alias)] = canonical
    return aliases


def load_positional_mapping() -> tuple[dict[int, dict[str, Any]] | None, str]:
    """Carga sólo un mapeo explícitamente validado y respaldado por evidencia."""

    selected = next((path for path in (MAPPING_LOCAL_PATH, MAPPING_EXAMPLE_PATH) if path.exists()), None)
    if selected is None:
        return None, "mapeo posicional SIO pendiente de validación; no existe archivo de mapeo"
    try:
        document = json.loads(selected.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        return None, f"mapeo posicional SIO pendiente de validación; archivo inválido ({exc.__class__.__name__})"
    if not isinstance(document, dict) or document.get("mapping_status") != "validado":
        return None, "mapeo posicional SIO pendiente de validación"
    evidence = document.get("evidence", [])
    evidence_items = [evidence] if isinstance(evidence, str) else evidence if isinstance(evidence, list) else []
    documented_evidence = any((ROOT / item).exists() if not Path(item).is_absolute() else Path(item).exists() for item in evidence_items if isinstance(item, str) and item.strip())
    raw_mapping = document.get("row_mapping")
    if not documented_evidence or not isinstance(raw_mapping, dict):
        return None, "mapeo posicional SIO pendiente de validación; falta evidencia documentada"
    mapping: dict[int, dict[str, Any]] = {}
    for position, field in raw_mapping.items():
        try:
            position_number = int(position)
        except (TypeError, ValueError):
            continue
        if position_number < 0:
            continue
        if isinstance(field, str) and field.strip():
            mapping[position_number] = {"target_field": field.strip()}
        elif isinstance(field, dict) and isinstance(field.get("target_field"), str) and field["target_field"].strip():
            mapping[position_number] = dict(field)
    if not mapping or not any(key(spec.get("target_field")) in {"fecha", "producto", "commodity", "precio"} for spec in mapping.values()):
        return None, "mapeo posicional SIO pendiente de validación; faltan campos mínimos"
    return mapping, f"mapeo posicional validado desde {selected.name}"


def map_positional_row(source: dict[str, Any], mapping: dict[int, dict[str, Any]]) -> dict[str, Any] | None:
    raw_row = next((value for name, value in source.items() if key(name) == "row"), None)
    if not isinstance(raw_row, (list, tuple)):
        return None
    mapped: dict[str, Any] = {}
    for position, specification in mapping.items():
        if position < len(raw_row):
            field = specification["target_field"].strip()
            mapped[field] = raw_row[position]
            source_label = text(specification.get("source_label"))
            if source_label:
                mapped[f"__source_label_{field}"] = source_label
            unit_field = text(specification.get("unit_field"))
            unit_value = text(specification.get("unit_value_if_label_matches"))
            if unit_field and unit_value and source_label:
                mapped[f"__unit_value_{unit_field}"] = unit_value
    return mapped or None


def normalize_commodity(value: Any, filename: str, aliases: dict[str, str]) -> tuple[str, str]:
    original = text(value)
    if original and key(original) in aliases:
        return aliases[key(original)], ""
    if original:
        return original, "commodity no catalogado; se conserva para revisión"
    filename_key = key(filename)
    for alias, canonical in sorted(aliases.items(), key=lambda item: len(item[0]), reverse=True):
        if alias and alias in filename_key:
            return canonical, "commodity inferido desde el nombre del archivo"
    return "Sin especificar", "commodity no identificado; revisar respuesta"


def value_for(row: dict[str, Any], *names: str) -> Any:
    normalized = row.get("__normalized_fields")
    if not isinstance(normalized, dict):
        normalized = {key(name): value for name, value in row.items() if name != "__normalized_fields"}
        row["__normalized_fields"] = normalized
    for name in names:
        if key(name) in normalized:
            return normalized[key(name)]
    return ""


def value_with_field(row: dict[str, Any], *names: str) -> tuple[Any, str]:
    normalized = row.get("__normalized_fields_with_names")
    if not isinstance(normalized, dict):
        normalized = {key(name): (value, str(name)) for name, value in row.items() if name not in {"__normalized_fields", "__normalized_fields_with_names"}}
        row["__normalized_fields_with_names"] = normalized
    for name in names:
        if key(name) in normalized:
            return normalized[key(name)]
    return "", ""


def explicit_unit_from_label(label: str, unit_value: str = "") -> str:
    if unit_value:
        return unit_value
    normalized = key(label)
    if "preciotn" in normalized or ("preciopor" in normalized and "tn" in normalized) or ("preciopor" in normalized and "tonelada" in normalized) or "$/tn" in label.lower():
        return "TN"
    return ""


def extract_records(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        result: list[dict[str, Any]] = []
        for item in payload:
            result.extend(extract_records(item))
        return result
    if isinstance(payload, dict):
        containers = {"data", "results", "records", "items", "rows", "operaciones", "d"}
        for name, value in payload.items():
            if str(name).lower() in containers:
                nested = extract_records(value)
                if nested:
                    return nested
        return [payload]
    return []


def read_matrix(values: list[tuple[Any, ...]]) -> tuple[list[dict[str, Any]], list[str]]:
    nonempty = [index for index, row in enumerate(values[:30]) if row and any(text(cell) for cell in row)]
    if not nonempty:
        return [], []
    header_index = max(nonempty, key=lambda index: sum(key(cell) in HEADER_KEYS for cell in values[index]))
    headers = [text(cell) or f"columna_{index + 1}" for index, cell in enumerate(values[header_index])]
    rows: list[dict[str, Any]] = []
    for values_row in values[header_index + 1:]:
        if values_row and any(text(cell) for cell in values_row):
            rows.append({headers[index]: values_row[index] if index < len(values_row) else "" for index in range(len(headers))})
    return rows, headers


class HTMLTableParser(HTMLParser):
    """Extrae tablas HTML simples sin interpretar formularios ni JavaScript."""

    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[tuple[Any, ...]]] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "table" and self._table is None:
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in {"th", "td"} and self._row is not None:
            self._cell = []

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"th", "td"} and self._cell is not None and self._row is not None:
            self._row.append("".join(self._cell).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            if self._row:
                self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            self.tables.append([tuple(row) for row in self._table])
            self._table = None


def text_encoding(path: Path) -> str:
    with path.open("rb") as handle:
        prefix = handle.read(4)
    if prefix.startswith(b"\xff\xfe"):
        return "utf-16-le"
    if prefix.startswith(b"\xfe\xff"):
        return "utf-16-be"
    if len(prefix) >= 2 and prefix[1] == 0:
        return "utf-16-le"
    if len(prefix) >= 2 and prefix[0] == 0:
        return "utf-16-be"
    return "utf-8-sig"


def read_file(path: Path) -> tuple[Iterable[dict[str, Any]], list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        encoding = text_encoding(path)
        with path.open("r", encoding=encoding, errors="replace", newline="") as handle:
            sample = handle.read(8192)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        with path.open("r", encoding=encoding, errors="replace", newline="") as handle:
            reader = csv.DictReader(handle, dialect=dialect)
            columns = [text(column) for column in (reader.fieldnames or [])]

        def rows() -> Iterable[dict[str, Any]]:
            with path.open("r", encoding=encoding, errors="replace", newline="") as handle:
                for row in csv.DictReader(handle, dialect=dialect):
                    yield dict(row)

        return rows(), columns
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(payload, dict) and isinstance(payload.get("d"), str):
            try:
                payload["d"] = json.loads(payload["d"])
            except json.JSONDecodeError:
                pass
        return extract_records(payload), []
    if suffix in {".html", ".htm"}:
        try:
            import pandas as pd
            frames = pd.read_html(path)
        except ImportError as exc:
            parser = HTMLTableParser()
            parser.feed(path.read_text(encoding="utf-8", errors="replace"))
            rows: list[dict[str, Any]] = []
            columns: list[str] = []
            for table in parser.tables:
                if not any(key(cell) in HEADER_KEYS for row in table[:5] for cell in row):
                    continue
                table_rows, table_columns = read_matrix(table)
                rows.extend(table_rows)
                columns.extend(column for column in table_columns if column not in columns)
            return rows, columns
        rows: list[dict[str, Any]] = []
        columns: list[str] = []
        for frame in frames:
            frame = frame.fillna("")
            table_rows, table_columns = read_matrix([tuple(frame.columns)] + [tuple(row) for row in frame.itertuples(index=False, name=None)])
            rows.extend(table_rows)
            columns.extend(column for column in table_columns if column not in columns)
        return rows, columns
    try:
        import pandas as pd
    except ImportError:
        pd = None
    if suffix == ".xls":
        if pd is None:
            raise RuntimeError("para leer XLS instale pandas y xlrd, o convierta a XLSX/CSV")
        frames = pd.read_excel(path, sheet_name=None, header=None)
        rows: list[dict[str, Any]] = []
        columns: list[str] = []
        for frame in frames.values():
            table_rows, table_columns = read_matrix([tuple(row) for row in frame.fillna("").itertuples(index=False, name=None)])
            rows.extend(table_rows)
            columns.extend(column for column in table_columns if column not in columns)
        return rows, columns
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("para leer XLSX instale openpyxl o convierta a CSV") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    columns: list[str] = []
    for sheet in workbook.worksheets:
        table_rows, table_columns = read_matrix([tuple(row) for row in sheet.values])
        rows.extend(table_rows)
        columns.extend(column for column in table_columns if column not in columns)
    return rows, columns


def first_text(source: dict[str, Any], *names: str) -> str:
    return text(value_for(source, *names))


def process_file(path: Path, aliases: dict[str, str], positional_mapping: dict[int, dict[str, Any]] | None, sample_pages: int, pagination_status: str = "no_aplica", row_sink: Any = None) -> tuple[list[dict[str, str]], dict[str, Any]]:
    source_rows, columns = read_file(path)
    rows: list[dict[str, str]] = []
    dates: list[date] = []
    prices: list[float] = []
    commodities: set[str] = set()
    positional_skipped = 0
    positional_applied = 0
    read_count = 0
    integrated_count = 0
    row_signatures: list[str] = []
    page_origin = page_number_from_path(path)
    sample_type = "exportacion_manual" if is_manual_export_file(path) else "paginacion_controlada" if page_origin else "piloto_una_pagina"
    for source in source_rows:
        read_count += 1
        used_positional = False
        source_id = first_text(source, "ID", "id_operacion_sio")
        has_positional_row = any(key(name) == "row" for name in source)
        raw_row = value_for(source, "Row") if has_positional_row else ""
        raw_row_signature = json.dumps(raw_row, ensure_ascii=False, sort_keys=True, separators=(",", ":")) if has_positional_row else json.dumps(source, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        if has_positional_row:
            if positional_mapping is None:
                positional_skipped += 1
                continue
            mapped_source = map_positional_row(source, positional_mapping)
            if mapped_source is None:
                positional_skipped += 1
                continue
            source = mapped_source
            if source_id:
                source["id_operacion_sio"] = source_id
            positional_applied += 1
            used_positional = True
        market_date = parse_date(value_for(source, "Fecha Operacion", "Fecha de Operacion", "Fecha Declaración", "Fecha Declaracion", "Fecha de Concertación", "Fecha de Concertacion", "Fecha de Entrega", "Fecha Concertación", "Fecha Concertacion", "Fecha"))
        raw_commodity = value_for(source, "Producto", "Grano", "Especie", "Commodity")
        commodity, commodity_note = normalize_commodity(raw_commodity, path.name, aliases)
        raw_price, detected_price_field = value_with_field(source, "Precio/TN Monto", "Precio/Monto", "Precio/TN", "Precio TN", "Precio unitario", "Precio", "Monto", "Precio hecho", "Precio Hecho", "Cotización", "Cotizacion", "Valor")
        price = parse_number(raw_price)
        if not market_date and price is None and not raw_commodity:
            continue
        if market_date and row_sink is None:
            dates.append(market_date)
        if price is not None and row_sink is None:
            prices.append(price)
        price_original = text(raw_price)
        price_label = text(source.get("__source_label_precio")) or detected_price_field
        price_type_original = first_text(source, "precio_tipo_original", "Precio tipo original", "Precio")
        raw_currency, detected_currency_field = value_with_field(source, "Precio/TN Moneda", "Moneda", "Currency", "Código moneda", "Codigo moneda")
        moneda, moneda_explicitamente_informada, moneda_inferida = extract_explicit_currency(price_original)
        currency_original = price_original if moneda_explicitamente_informada == "sí" else text(raw_currency)
        currency_label = price_label if moneda_explicitamente_informada == "sí" else detected_currency_field
        if moneda_explicitamente_informada != "sí" and text(raw_currency):
            normalized_currency = normalize_explicit_currency_field(raw_currency)
            if normalized_currency in {"ARS", "USD"}:
                moneda = normalized_currency
                moneda_explicitamente_informada = "sí"
        source_unit = first_text(source, "Unidad", "Unit")
        explicit_price_unit = explicit_unit_from_label(price_label, text(source.get("__unit_value_unidad")))
        unidad = source_unit or explicit_price_unit
        precio_unidad = price if explicit_price_unit and price is not None else None
        raw_total, detected_total_field = value_with_field(source, "Precio total", "Monto total", "Total")
        precio_total = parse_number(raw_total)
        # En GetOperaciones la columna "Tipo" describe la operación comercial
        # (Compraventa/Canje), mientras que la columna siguiente contiene el
        # tipo técnico de precio (por ejemplo, "Precio Hecho"). Para el
        # dashboard, la primera es la dimensión operativa visible y la segunda
        # se conserva en precio_tipo_original como metadata.
        operation_type = first_text(source, "tipo_operacion", "Tipo de operación", "Tipo de Operacion")
        tipo = operation_type or first_text(source, "tipo_precio", "Tipo de precio", "Tipo Precio", "Price Type", "Tipo") or price_type_original
        volumen, detected_volume_field = value_with_field(source, "Cant. (TN)", "Cant TN", "Cantidad (TN)", "Cantidad TN", "Volumen TN", "Toneladas", "Volumen", "Cantidad", "TN")
        volume_label = text(source.get("__source_label_volumen")) or detected_volume_field
        volumen_unidad = first_text(source, "volumen_unidad", "Unidad de volumen", "Unidad volumen", "Volume Unit") or text(source.get("__unit_value_volumen_unidad"))
        if not volumen_unidad and key(volume_label) in {"canttn", "cantidadtn", "cantidadtns", "volumentn", "toneladas", "tn"}:
            volumen_unidad = "TN"
        row_missing: list[str] = []
        if moneda_explicitamente_informada == "sí":
            row_missing.append(f"moneda extraída explícitamente del campo precio original: {moneda}" if currency_label == price_label else f"moneda informada explícitamente en campo original: {moneda}")
        else:
            moneda = "Sin especificar"
            row_missing.append("moneda no informada explícitamente en endpoint GetOperaciones; no se infiere del valor original")
        if not unidad:
            unidad = "Sin especificar"
            row_missing.append("falta unidad")
        if not tipo:
            tipo = "Sin especificar"
            row_missing.append("falta tipo_precio")
        if not market_date:
            row_missing.append("falta fecha válida")
        if price is None:
            row_missing.append("falta precio válido")
        operation = first_text(source, "Operación", "Operacion", "operacion")
        payment = first_text(source, "condicion_pago", "Condición de Pago", "Condicion de Pago", "Pago")
        commercial = first_text(source, "Condición comercial", "Condicion comercial", "Condición", "Condicion", "Entrega")
        observation = first_text(source, "Observación", "Observaciones", "Nota", "Notas")
        zero_type = classify_zero_price(price_original, price, operation)
        zero_flag = "sí" if price == 0 else "no"
        price_valid_for_series = "sí" if price is not None and price > 0 else "no"
        pilot_note = "integración piloto una página GetOperaciones" if used_positional else ""
        zero_note = f"precio cero clasificado como {zero_type}; excluido de series de precios" if zero_flag == "sí" else ""
        notes = "; ".join(dict.fromkeys([part for part in [observation, commodity_note, pilot_note, zero_note] + row_missing if part]))
        source_name = first_text(source, "Fuente", "Source") or DEFAULT_SOURCE
        delivery_place = first_text(source, "Zona", "Lugar Entrega", "Lugar de entrega", "lugar_entrega")
        pilot_status = "sí" if market_date and commodity != "Sin especificar" and price is not None and source_name and price_label and price_label != "Sin especificar" and explicit_price_unit else "no"
        dashboard_status = "parcial_piloto" if pilot_status == "sí" and moneda_explicitamente_informada == "sí" and explicit_price_unit and used_positional else "no"
        commodities.add(commodity)
        mapped_row = {
            "fecha": market_date.isoformat() if market_date else "",
            "año": str(market_date.year) if market_date else "",
            "mes": str(market_date.month) if market_date else "",
            "commodity": commodity,
            "fuente": source_name,
            "mercado": first_text(source, "Mercado", "Market"),
            "tipo_precio": tipo,
            "precio_tipo_original": price_type_original,
            "precio_unidad": "" if precio_unidad is None else f"{precio_unidad:g}",
            "precio_total": "" if precio_total is None else f"{precio_total:g}",
            "campo_precio_original": price_label or "Sin especificar",
            "valor_precio_original": price_original,
            "precio_original_texto": price_original,
            "moneda": moneda,
            "moneda_explicitamente_informada": moneda_explicitamente_informada,
            "moneda_inferida": moneda_inferida,
            "campo_moneda_original": currency_label or "Sin especificar",
            "valor_moneda_original": currency_original,
            "unidad": unidad,
            "precio": "" if price is None else f"{price:g}",
            "volumen": volumen,
            "volumen_unidad": volumen_unidad,
            "campo_volumen_original": volume_label or "Sin especificar",
            "procedencia": first_text(source, "Procedencia", "Procedencia Pcia", "Procedencia Localid."),
            "provincia": first_text(source, "Provincia", "Pcia", "Procedencia Pcia"),
            "localidad": first_text(source, "Localidad"),
            "zona": delivery_place,
            "lugar_entrega": delivery_place,
            "precio_puesto_en": first_text(source, "Precio puesto en", "Destino", "Puerto"),
            "operacion": operation,
            "condicion_pago": payment,
            "condicion_comercial": commercial,
            "frecuencia": first_text(source, "Frecuencia", "Frequency"),
            "archivo_origen": path.name,
            "fecha_integracion": date.today().isoformat(),
            "observaciones": notes,
            "apto_piloto": pilot_status,
            "apto_dashboard": dashboard_status,
            "pagina_origen": page_origin,
            "id_operacion_sio": source_id or first_text(source, "id_operacion_sio", "ID"),
            "muestra_tipo": sample_type,
            "muestra_paginas": str(sample_pages),
            "estado_paginacion": pagination_status if sample_type == "paginacion_controlada" else "no_aplica",
            "precio_cero_flag": zero_flag,
            "precio_cero_tipo": zero_type,
            "precio_valido_para_serie": price_valid_for_series,
            "_row_signature": raw_row_signature,
        }
        if row_sink is None:
            rows.append(mapped_row)
        else:
            row_sink(mapped_row)
        integrated_count += 1
        if row_sink is None:
            row_signatures.append(raw_row_signature)
    return rows, {"read": read_count, "integrated": integrated_count, "positional_skipped": positional_skipped, "positional_applied": positional_applied, "columns": columns, "commodities": sorted(commodities), "dates": dates, "prices": prices, "page_origin": page_origin, "row_signatures": row_signatures}


def real_files() -> list[Path]:
    if not RAW_DIR.exists():
        return []
    return sorted(path for path in RAW_DIR.iterdir() if path.is_file() and path.suffix.lower() in EXTENSIONS and not any(marker in path.stem.lower() for marker in NON_REAL_MARKERS))


def is_latest_snapshot_file(path: Path) -> bool:
    return bool(LATEST_SNAPSHOT_PATTERN.fullmatch(path.name))


def snapshot_capture_timestamp(path: Path) -> str:
    match = re.search(r"_(\d{8})_(\d{6})\.json$", path.name, flags=re.I)
    if match:
        try:
            return datetime.strptime(f"{match.group(1)}{match.group(2)}", "%Y%m%d%H%M%S").isoformat(timespec="seconds")
        except ValueError:
            pass
    return datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")


def snapshot_identity(row: dict[str, str]) -> tuple[str, ...]:
    operation_id = text(row.get("id_operacion_sio"))
    if operation_id:
        return ("id", operation_id)
    return (
        "fallback",
        text(row.get("fecha")),
        text(row.get("commodity")),
        text(row.get("precio_original_texto")),
        text(row.get("volumen")),
        text(row.get("procedencia")),
        text(row.get("lugar_entrega")),
        text(row.get("condicion_comercial")),
    )


def deduplicate_snapshot_rows(rows: list[dict[str, str]]) -> tuple[list[dict[str, str]], int, int]:
    """Deduplica snapshots por ID y, si falta, por la clave estable documentada."""

    kept: list[dict[str, str]] = []
    seen: dict[tuple[str, ...], dict[str, str]] = {}
    duplicates = 0
    conflicts = 0
    for row in rows:
        identity = snapshot_identity(row)
        previous = seen.get(identity)
        if previous is not None:
            duplicates += 1
            if raw_row_signature(previous) != raw_row_signature(row):
                conflicts += 1
            continue
        seen[identity] = row
        kept.append(row)
    return kept, duplicates, conflicts


def write_snapshot_output(path: Path, rows: list[dict[str, str]]) -> None:
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=SNAPSHOT_COLUMNS, delimiter=";", extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def read_snapshot_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle, delimiter=";")]


def light_history_row(row: dict[str, str], dashboard_date: str = "") -> dict[str, str]:
    """Reduce una operación SIO a la memoria persistente que puede versionarse."""

    result = {column: text(row.get(column)) for column in LIGHT_HISTORY_COLUMNS}
    operation_type = text(row.get("tipo_operacion"))
    if operation_type in {"Compraventa", "Canje"}:
        result["tipo_precio"] = operation_type
    result["fecha_actualizacion_dashboard"] = result["fecha_actualizacion_dashboard"] or dashboard_date or date.today().isoformat()
    return result


def raw_operation_type_map() -> dict[str, str]:
    """Lee el tipo comercial de snapshots raw locales sin publicar esos archivos."""

    operation_types: dict[str, str] = {}
    for path in sorted(RAW_DIR.glob("SIO_latest_GetOperaciones_*.json")):
        try:
            document = json.loads(path.read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError):
            continue
        items = document.get("d", {}).get("Items", []) if isinstance(document, dict) else []
        for item in items if isinstance(items, list) else []:
            if not isinstance(item, dict):
                continue
            raw_row = item.get("Row", [])
            if not isinstance(raw_row, list) or len(raw_row) <= 4:
                continue
            operation_id = text(item.get("ID") or raw_row[0])
            operation_type = text(raw_row[4])
            if operation_id and operation_type in {"Compraventa", "Canje"}:
                operation_types[operation_id] = operation_type
    return operation_types


def apply_raw_operation_types(rows: list[dict[str, str]]) -> None:
    """Completa snapshots antiguos cuando el campo operativo quedó fuera del CSV liviano."""

    operation_types = raw_operation_type_map()
    if not operation_types:
        return
    for row in rows:
        operation_type = text(row.get("tipo_operacion")) or operation_types.get(text(row.get("id_operacion_sio")), "")
        if operation_type in {"Compraventa", "Canje"}:
            row["tipo_precio"] = operation_type


def write_light_history(path: Path, rows: list[dict[str, str]]) -> None:
    DASHBOARD_DIR.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=LIGHT_HISTORY_COLUMNS, delimiter=";", extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(light_history_row(row) for row in rows)


def integrate_latest_snapshot(path: Path, aliases: dict[str, str], positional_mapping: dict[int, dict[str, Any]] | None) -> dict[str, Any]:
    """Integra sólo el último JSON latest y lo suma al histórico de snapshots."""

    mapping = positional_mapping or LATEST_POSITIONAL_MAPPING
    rows, diagnostics = process_file(path, aliases, mapping, 1, "no_aplica")
    capture_timestamp = snapshot_capture_timestamp(path)
    for row in rows:
        row["fecha_descarga_snapshot"] = capture_timestamp
        row["frecuencia"] = text(row.get("frecuencia")) or "diaria"
        row["muestra_tipo"] = "snapshot_latest"
        row["observaciones"] = "; ".join(dict.fromkeys(filter(None, [text(row.get("observaciones")), "snapshot de últimas operaciones; no es histórico completo"])))
    latest_rows, latest_duplicates, latest_conflicts = deduplicate_snapshot_rows(rows)
    if not latest_rows:
        history_rows = read_snapshot_rows(HISTORIC_SNAPSHOTS_OUTPUT_PATH) or read_snapshot_rows(LIGHT_HISTORY_OUTPUT_PATH)
        return {"rows": [], "diagnostics": diagnostics, "latest_duplicates": latest_duplicates, "latest_conflicts": latest_conflicts, "history_rows": len(history_rows), "new_rows": 0, "history_duplicates": 0, "capture_timestamp": capture_timestamp}

    processed_history = read_snapshot_rows(HISTORIC_SNAPSHOTS_OUTPUT_PATH)
    versioned_history = read_snapshot_rows(LIGHT_HISTORY_OUTPUT_PATH)
    history_before, _, _ = deduplicate_snapshot_rows(processed_history + versioned_history)
    existing_ids = {snapshot_identity(row) for row in history_before}
    new_rows = sum(1 for row in latest_rows if snapshot_identity(row) not in existing_ids)
    combined_rows, history_duplicates, history_conflicts = deduplicate_snapshot_rows(history_before + latest_rows)
    apply_raw_operation_types(combined_rows)
    latest_by_identity = {snapshot_identity(row): row for row in latest_rows}
    for row in combined_rows:
        latest_row = latest_by_identity.get(snapshot_identity(row))
        if latest_row is not None:
            # Se conserva la primera captura para poder distinguir operaciones
            # nuevas de repetidas en el reporte diario. La captura más reciente
            # se obtiene de LATEST_SNAPSHOT.csv al preparar el dashboard.
            row["fecha_descarga_snapshot"] = text(row.get("fecha_descarga_snapshot")) or capture_timestamp
            row["observaciones"] = "; ".join(dict.fromkeys(filter(None, [text(row.get("observaciones")), "observada en snapshot latest",])))
            row["fecha_actualizacion_dashboard"] = date.today().isoformat()
    write_snapshot_output(LATEST_SNAPSHOT_OUTPUT_PATH, latest_rows)
    write_snapshot_output(HISTORIC_SNAPSHOTS_OUTPUT_PATH, combined_rows)
    write_light_history(LIGHT_HISTORY_OUTPUT_PATH, combined_rows)
    return {
        "rows": latest_rows,
        "diagnostics": diagnostics,
        "latest_duplicates": latest_duplicates,
        "latest_conflicts": latest_conflicts,
        "history_rows": len(combined_rows),
        "new_rows": new_rows,
        "history_duplicates": history_duplicates,
        "history_conflicts": history_conflicts,
        "capture_timestamp": capture_timestamp,
    }


def row_signature(row: dict[str, str]) -> tuple[str, ...]:
    excluded = {"archivo_origen", "fecha_integracion", "observaciones", "pagina_origen", "muestra_paginas", "muestra_tipo", "estado_paginacion"}
    return tuple(str(row.get(column, "")) for column in OUTPUT_COLUMNS if column not in excluded)


def raw_row_signature(row: dict[str, str]) -> str:
    return text(row.get("_row_signature")) or json.dumps(row_signature(row), ensure_ascii=False, separators=(",", ":"))


def deduplicate_rows(rows: list[dict[str, str]]) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]]]:
    kept: list[dict[str, str]] = []
    duplicates: list[dict[str, str]] = []
    conflicts: list[dict[str, str]] = []
    seen_by_id: dict[str, dict[str, str]] = {}
    seen_by_key: dict[tuple[str, ...], dict[str, str]] = {}
    for row in rows:
        operation_id = text(row.get("id_operacion_sio"))
        key_value = (text(row.get("fecha")), text(row.get("commodity")), text(row.get("precio")), text(row.get("moneda")), text(row.get("unidad")), text(row.get("volumen")), text(row.get("procedencia")), text(row.get("lugar_entrega")), text(row.get("tipo_precio")))
        previous = seen_by_id.get(operation_id) if operation_id else seen_by_key.get(key_value)
        if previous is None:
            kept.append(row)
            if operation_id:
                seen_by_id[operation_id] = row
            else:
                seen_by_key[key_value] = row
            continue
        if raw_row_signature(previous) == raw_row_signature(row):
            duplicate_type = "duplicado por id_operacion_sio y Row" if operation_id else "duplicado exacto por Row"
            duplicates.append({"id_operacion_sio": operation_id, "archivo_origen": row.get("archivo_origen", ""), "pagina_origen": row.get("pagina_origen", ""), "tipo": duplicate_type, "motivo": "duplicado exacto; se conserva la primera aparición"})
        else:
            conflicts.append({"id_operacion_sio": operation_id, "archivo_origen": row.get("archivo_origen", ""), "pagina_origen": row.get("pagina_origen", ""), "tipo": "conflicto de ID/claves", "motivo": "mismo ID/claves pero Row o contenido diferente; se conservan ambas filas para revisión"})
            kept.append(row)
    return kept, duplicates, conflicts


def update_paginated_report(files: list[Path], diagnostics: list[dict[str, Any]], rows: list[dict[str, str]], duplicates: list[dict[str, str]], conflicts: list[dict[str, str]], page_count: int, pagination_status: str) -> None:
    report_paths = [path for path in (PAGINATED_REPORT_PATH, PAGINATION_REPORT_PATH) if path.exists()]
    if not report_paths:
        return
    total_read = sum(int(item.get("read", 0)) for item in diagnostics)
    columns = [column for column in OUTPUT_COLUMNS if any(row.get(column) for row in rows)]
    duplicate_by_id = sum(1 for item in duplicates if item.get("id_operacion_sio"))
    duplicate_by_row = sum(1 for item in duplicates if "Row" in item.get("tipo", ""))
    warning = "paginación no validada; páginas repetidas" if pagination_status == "duplicada" else "sin advertencia de repetición completa de páginas"
    integration_section = "\n".join([
        "## Resultado de integración", "", f"- Filas leídas antes de deduplicar: {total_read}.", f"- Archivos procesados: {len(files)}.", f"- Páginas procesadas: {page_count}.", f"- Duplicados exactos eliminados: {len(duplicates)}; con ID: {duplicate_by_id}; por Row: {duplicate_by_row}.", f"- Conflictos conservados para revisión: {len(conflicts)}.", f"- Filas finales: {len(rows)}.", f"- Salida técnica: `data/commodities_sio/processed/{PAGINATED_OUTPUT_PATH.name}`; no reemplaza `COMMODITIES_SIO_INTEGRADO.csv`.", f"- Estado derivado de paginación: {pagination_status}.", f"- Observación: {warning}.", f"- Columnas principales: {', '.join(columns)}.", "",
    ])
    for report_path in report_paths:
        report = report_path.read_text(encoding="utf-8")
        report = re.sub(r"## Resultado de integración\n.*?(?=\n## (?:Resultado de auditoría|Paginación y duplicados))", integration_section.rstrip(), report, flags=re.S)
        report = report.replace("## Paginación y duplicados\n\nPendiente de ejecutar `auditar_commodities_sio.py`.", f"## Paginación y duplicados\n\n- Estado derivado en integración: `{pagination_status}`.\n- Observación: {warning}.")
        report_path.write_text(report, encoding="utf-8")


def write_output(path: Path, rows: list[dict[str, str]]) -> None:
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS, delimiter=";", extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def is_paginated_file(path: Path) -> bool:
    return page_number_from_path(path) != ""


def page_number_from_path(path: Path) -> str:
    match = re.search(r"(?:page[_-]|observed_pcurrentpage[_-])(\d+)", path.stem, flags=re.I)
    return match.group(1) if match else ""


def is_observed_pagination_file(path: Path) -> bool:
    return bool(re.search(r"observed_pcurrentpage[_-]\d+", path.stem, flags=re.I))


def is_manual_export_file(path: Path) -> bool:
    return bool(re.fullmatch(r"SIO_exportar_operaciones_.*", path.stem, flags=re.I)) and path.suffix.lower() in {".xlsx", ".xls", ".csv"}


def process_group(files: list[Path], aliases: dict[str, str], positional_mapping: dict[int, dict[str, Any]] | None, sample_pages: int, pagination_status: str) -> tuple[list[dict[str, str]], list[dict[str, Any]], int]:
    rows_all: list[dict[str, str]] = []
    diagnostics_all: list[dict[str, Any]] = []
    errors = 0
    for path in files:
        try:
            rows, diagnostics = process_file(path, aliases, positional_mapping, sample_pages, pagination_status)
            rows_all.extend(rows)
            diagnostics["archivo_origen"] = path.name
            diagnostics_all.append(diagnostics)
            print(f"Archivo procesado: {path.name}")
            print(f"  Filas leídas: {diagnostics['read']}; filas integradas: {diagnostics['integrated']}")
            if diagnostics["positional_skipped"]:
                print(f"  Filas Row posicionales omitidas: {diagnostics['positional_skipped']}")
            if diagnostics["positional_applied"]:
                print(f"  Filas Row con mapeo validado aplicado: {diagnostics['positional_applied']}")
        except Exception as exc:
            errors += 1
            print(f"ERROR en {path.name}: {exc}")
            print("  Sugerencia: revisar encabezados o conservar la respuesta original para ajustar el mapeo.")
    return rows_all, diagnostics_all, errors


def process_manual_exports_streaming(files: list[Path], aliases: dict[str, str], positional_mapping: dict[int, dict[str, Any]] | None) -> tuple[int, list[dict[str, Any]], int, int, int, int]:
    """Integra exportaciones grandes sin conservar cientos de miles de filas en memoria."""

    if not files:
        return 0, [], 0, 0, 0, 0
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    temporary_path = MANUAL_EXPORT_OUTPUT_PATH.with_suffix(".csv.tmp")
    seen_by_id: dict[str, str] = {}
    seen_by_key: dict[tuple[str, ...], str] = {}
    diagnostics_all: list[dict[str, Any]] = []
    counts = {"written": 0, "duplicates": 0, "conflicts": 0}

    def consume(row: dict[str, str]) -> None:
        operation_id = text(row.get("id_operacion_sio"))
        key_value = (text(row.get("fecha")), text(row.get("commodity")), text(row.get("precio")), text(row.get("moneda")), text(row.get("unidad")), text(row.get("volumen")), text(row.get("procedencia")), text(row.get("lugar_entrega")), text(row.get("tipo_precio")))
        signature = raw_row_signature(row)
        previous = seen_by_id.get(operation_id) if operation_id else seen_by_key.get(key_value)
        if previous is not None:
            if previous == signature:
                counts["duplicates"] += 1
                return
            counts["conflicts"] += 1
        if operation_id:
            seen_by_id.setdefault(operation_id, signature)
        else:
            seen_by_key.setdefault(key_value, signature)
        writer.writerow(row)
        counts["written"] += 1

    try:
        with temporary_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS, delimiter=";", extrasaction="ignore", lineterminator="\n")
            writer.writeheader()
            for path in files:
                try:
                    _, diagnostics = process_file(path, aliases, positional_mapping, 1, "no_aplica", row_sink=consume)
                    diagnostics["archivo_origen"] = path.name
                    diagnostics_all.append(diagnostics)
                    print(f"Archivo procesado: {path.name}")
                    print(f"  Filas leídas: {diagnostics['read']}; filas integradas: {diagnostics['integrated']}")
                except Exception as exc:
                    print(f"ERROR en {path.name}: {exc}")
                    return 0, diagnostics_all, 1, counts["duplicates"], counts["conflicts"], counts["written"]
        temporary_path.replace(MANUAL_EXPORT_OUTPUT_PATH)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    return 0, diagnostics_all, 0, counts["duplicates"], counts["conflicts"], counts["written"]


def main() -> int:
    files = real_files()
    if not files:
        print("No hay archivos reales de SIO Granos en data/commodities_sio/raw/.")
        print("Se preservan las salidas procesadas existentes; no se generan archivos vacíos.")
        return 0
    aliases = read_aliases()
    positional_mapping, mapping_status = load_positional_mapping()
    print(f"Mapeo posicional: {mapping_status}")
    page_files = [path for path in files if is_paginated_file(path)]
    latest_snapshot_files = [path for path in files if is_latest_snapshot_file(path)]
    latest_snapshot_file = latest_snapshot_files[-1] if latest_snapshot_files else None
    manual_export_files = [path for path in files if is_manual_export_file(path)] if not latest_snapshot_files else []
    base_files = [path for path in files if not is_paginated_file(path) and not is_manual_export_file(path) and not is_latest_snapshot_file(path)]
    observed_page_files = [path for path in page_files if is_observed_pagination_file(path)]
    status_page_files = observed_page_files or page_files
    page_numbers = {int(page_number_from_path(path)) for path in status_page_files if page_number_from_path(path)}
    sample_pages = len(page_numbers) or 1
    latest_error = 0
    if latest_snapshot_file:
        try:
            latest_result = integrate_latest_snapshot(latest_snapshot_file, aliases, positional_mapping)
            latest_rows = latest_result["rows"]
            if latest_rows:
                print(f"Snapshot latest integrado: {len(latest_rows)} operaciones válidas en {LATEST_SNAPSHOT_OUTPUT_PATH}")
                print(f"Histórico de snapshots: {latest_result['history_rows']} filas acumuladas en {HISTORIC_SNAPSHOTS_OUTPUT_PATH}")
                print(f"Operaciones nuevas acumuladas: {latest_result['new_rows']}; duplicados omitidos: {latest_result['history_duplicates']}; conflictos omitidos: {latest_result['history_conflicts']}")
            else:
                latest_error = 1
                print("ERROR: el snapshot latest no produjo operaciones integrables; no se reemplazaron las salidas de snapshots.")
                print(f"  Filas leídas: {latest_result['diagnostics']['read']}; filas omitidas por mapeo: {latest_result['diagnostics']['positional_skipped']}")
        except Exception as exc:
            latest_error = 1
            print(f"ERROR integrando snapshot latest {latest_snapshot_file.name}: {exc}")
    technical_rows_unchecked, page_diagnostics, page_errors = process_group(page_files, aliases, positional_mapping, sample_pages, "no_probada")
    status_file_names = {path.name for path in status_page_files}
    status_diagnostics = [item for item in page_diagnostics if item.get("archivo_origen") in status_file_names]
    page_signature_groups = {str(item["page_origin"]): item.get("row_signatures", []) for item in status_diagnostics if item.get("row_signatures")}
    first_page_signatures = next(iter(page_signature_groups.values()), [])
    pagination_status = "no_probada"
    if len(page_signature_groups) >= 2:
        pagination_status = "duplicada" if first_page_signatures and all(signatures == first_page_signatures for signatures in page_signature_groups.values()) else "parcial"
    elif page_signature_groups:
        pagination_status = "parcial"
    base_rows_unchecked, base_diagnostics, base_errors = process_group(base_files, aliases, positional_mapping, 1, "no_aplica")
    _, manual_diagnostics, manual_errors, manual_duplicates_count, manual_conflicts_count, manual_count = process_manual_exports_streaming(manual_export_files, aliases, positional_mapping)
    base_rows, base_duplicates, base_conflicts = deduplicate_rows(base_rows_unchecked)
    technical_rows, duplicates, conflicts = deduplicate_rows(technical_rows_unchecked)
    if pagination_status == "duplicada":
        for row in technical_rows:
            row["estado_paginacion"] = "duplicada"
            note = "paginación no validada; páginas repetidas"
            row["observaciones"] = "; ".join(dict.fromkeys(filter(None, [row.get("observaciones", ""), note])))
    elif technical_rows:
        for row in technical_rows:
            row["estado_paginacion"] = pagination_status
    if base_rows:
        write_output(OUTPUT_PATH, base_rows)
        print(f"Integración piloto principal: {len(base_rows)} filas en {OUTPUT_PATH}")
    else:
        print("No se encontraron filas piloto base; se preserva COMMODITIES_SIO_INTEGRADO.csv existente.")
    if technical_rows:
        write_output(PAGINATED_OUTPUT_PATH, technical_rows)
        print(f"Muestra paginada técnica: {len(technical_rows)} filas en {PAGINATED_OUTPUT_PATH}")
    if manual_count:
        print(f"Exportación manual SIO: {manual_count} filas en {MANUAL_EXPORT_OUTPUT_PATH}")
    elif manual_export_files:
        print("No se integraron filas de exportación manual; se preserva COMMODITIES_SIO_EXPORTACION_MANUAL.csv existente.")
    elif latest_snapshot_files and any(is_manual_export_file(path) for path in files):
        print("La exportación manual permanece fuera de esta corrida diaria de snapshots y no se modifica.")
    update_paginated_report(page_files, page_diagnostics, technical_rows, duplicates, conflicts, sample_pages, pagination_status)
    errors = latest_error + page_errors + base_errors + manual_errors
    print(f"Piloto base: duplicados exactos={len(base_duplicates)}; conflictos={len(base_conflicts)}")
    print(f"Duplicados exactos eliminados: {len(duplicates)}; conflictos conservados: {len(conflicts)}")
    if manual_export_files:
        print(f"Exportación manual: duplicados exactos={manual_duplicates_count}; conflictos={manual_conflicts_count}")
    print(f"Estado de paginación: {pagination_status}.")
    if pagination_status == "duplicada":
        print("Advertencia: paginación no validada; páginas repetidas. Se conservan sólo registros únicos con trazabilidad.")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
