import io
import unittest
from datetime import date

from openpyxl import Workbook
from scripts.carnes.magyp_workbooks import parse_workbook
from scripts.carnes.model import DataError

ASOF=date(2026,10,22)
OLD=['Buey','Novillito','Novillo','Ternera','Ternero','Torito','Toro','Vaquillona','Vaca']
MODERN=[('M.E.J.',['2 dientes']),('Novillito',['2 dientes','4 dientes']),
    ('Novillo',['6 dientes','8 dientes','Más de 8 dientes']),
    ('Toro',['4 dientes','6 dientes','8 dientes','Más de 8 dientes']),
    ('Vaquillona',['2 dientes','4 dientes']),('Vaca',['6 dientes','8 dientes','Más de 8 dientes'])]


def encoded(book):
    stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()


def category(sheet,year,start=1,family='slaughter'):
    measure={'slaughter':'FAENA DE BOVINOS - EN CABEZAS',
        'production':'PRODUCCIÓN DE CARNE BOVINA - EN TONELADAS RES CON HUESO',
        'carcass_weight':'PESO MEDIO DE LA RES BOVINA - EN KG RES CON HUESO'}[family]
    sheet.cell(start,1,f'AÑO {year} - {measure}')
    header=start+1;sheet.cell(header,1,'Mes /Categoría y Edad')
    groups=[]
    if year<=2019:groups.extend((name,[]) for name in OLD)
    if year>=2019:groups.extend(MODERN)
    groups.append(('Total General',[]));column=2
    for name,ages in groups:
        sheet.cell(header,column,name)
        if len(ages)>1:sheet.merge_cells(start_row=header,start_column=column,end_row=header,end_column=column+len(ages)-1)
        for offset,age in enumerate(ages):sheet.cell(header+1,column+offset,age)
        for offset in range(max(1,len(ages))):
            sheet.cell(header+2,column+offset,100.25)
            sheet.cell(header+3,column+offset,101.5)
            sheet.cell(header+4,column+offset,201.75)
        column+=max(1,len(ages))
    sheet.cell(header+2,1,'ENERO');sheet.cell(header+3,1,'FEBRERO')
    sheet.cell(header+4,1,'TOTAL POR EDAD' if year>=2019 else 'TOTAL POR CATEGORÍA')
    return header


def slaughter(year=2026):
    book=Workbook();cover=book.active;cover.title='Presentación'
    cover['A1']='Especies: Bovinos y Bubalinos' if year<=2019 else 'Especie: Bovinos'
    for name,metadata in [
        ('POR PROVINCIA',['Provincia']),
        ('POR ESTABLECIMIENTO',['Nº Est.','CUIT','Razón Social Frigorifico','Provincia']),
        ('POR USUARIO DE FAENA',['CUIT','Razón Social Usuario','Actividad'])]:
        sheet=book.create_sheet(name);sheet.cell(1,1,f'AÑO {year} - FAENA DE BOVINOS - EN CABEZAS')
        sheet.append(metadata+['ENERO','FEBRERO',f'ACUMULADO {year}'])
        identity={'POR PROVINCIA':['CORRIENTES'],'POR ESTABLECIMIENTO':[69,30540080298,'Frigorífico de prueba','CORRIENTES'],
            'POR USUARIO DE FAENA':[30540080298,'Usuario de prueba','Matarife Abastecedor']}[name]
        sheet.append(identity+[10.25,0,10.25])
    category(book.create_sheet('POR CATEGORIA' if year==2018 else 'POR CATEGORÍA Y EDAD'),year)
    return book


def production(year=2026):
    book=Workbook();sheet=book.active;sheet.title='PRODUCCIÓN POR CATEGORÍA Y EDAD'
    for start,family in [(1,'slaughter'),(10,'production'),(19,'carcass_weight')]:category(sheet,year,start,family)
    sheet['A30']='Datos provisorios sujetos a modificación'
    return book


class MagypWorkbookTests(unittest.TestCase):
    def test_fractional_heads_entity_ids_activity_and_source_cells_are_preserved(self):
        rows=parse_workbook(encoded(slaughter()),'slaughter_detail',ASOF,2026)
        establishment=next(r for r in rows if r['entity_type']=='establishment' and r['observation_date']=='2026-01-01')
        self.assertEqual(establishment['value'],'10.25')
        self.assertEqual(establishment['entity_id'],'69')
        self.assertEqual(establishment['entity_tax_id'],'30540080298')
        self.assertEqual(establishment['source_locator'],'POR ESTABLECIMIENTO!E3')
        self.assertIn('fractional_published_value',establishment['quality_flags'])
        user=next(r for r in rows if r['entity_type']=='slaughter_user')
        self.assertEqual(user['entity_activity'],'Matarife Abastecedor')
        self.assertTrue(any(r['value']=='0' and r['value_status']=='observed' for r in rows))

    def test_2018_combined_species_old_categories_do_not_become_dentition(self):
        rows=parse_workbook(encoded(slaughter(2018)),'slaughter_detail',ASOF,2018)
        row=next(r for r in rows if r['category']=='Ternera' and r['frequency']=='monthly')
        self.assertEqual(row['species'],'bovinos_y_bubalinos')
        self.assertEqual(row['age_class'],'')
        self.assertEqual(row['methodology'],'pre_2019_weight_sex_categories')
        self.assertIn('combined_species_not_separable',row['quality_flags'])

    def test_2019_classifications_are_distinct_even_when_category_name_matches(self):
        rows=parse_workbook(encoded(slaughter(2019)),'slaughter_detail',ASOF,2019)
        january=[r for r in rows if r['observation_date']=='2019-01-01' and r['frequency']=='monthly' and r['category']=='Novillito']
        self.assertEqual(len(january),3)
        self.assertEqual({r['methodology'] for r in january},{'pre_2019_weight_sex_categories','resolution_32_2018_dentition'})
        self.assertTrue(all('classification_transition_2019_no_bridge' in r['quality_flags'] for r in january))

    def test_production_and_weight_are_separate_published_measures(self):
        rows=parse_workbook(encoded(production()),'production_detail',ASOF,2026)
        self.assertEqual({r['family'] for r in rows},{'slaughter','production','carcass_weight'})
        units={family:{r['unit'] for r in rows if r['family']==family} for family in ['slaughter','production','carcass_weight']}
        self.assertEqual(units,{'slaughter':{'heads'},'production':{'tonnes_carcass_bone_in'},'carcass_weight':{'kg_carcass_bone_in'}})
        self.assertTrue(all(r['currency']=='NOT_APPLICABLE' for r in rows))
        self.assertTrue(all(r['weight_basis']=='carcass_bone_in' for r in rows if r['family']!='slaughter'))
        self.assertTrue(all('provisional_source_statement' in r['quality_flags'] for r in rows))

    def test_cumulative_period_has_end_and_cannot_be_summed_with_monthly(self):
        rows=parse_workbook(encoded(slaughter()),'slaughter_detail',ASOF,2026)
        total=next(r for r in rows if r['entity_type']=='province' and r['frequency']=='year_to_date')
        self.assertEqual(total['observation_date'],'2026-01-01')
        self.assertEqual(total['period_end'],'2026-02-28')
        self.assertIn('published_rollup',total['quality_flags'])
        self.assertIn('published_cumulative',total['quality_flags'])
        all_categories=[r for r in rows if r['category']=='Total general']
        self.assertTrue(all('published_rollup' in r['quality_flags'] for r in all_categories))

    def test_formula_without_cached_result_remains_missing_not_zero(self):
        book=slaughter();sheet=book['POR PROVINCIA'];sheet['B3']='=SUM(1,2)'
        rows=parse_workbook(encoded(book),'slaughter_detail',ASOF,2026)
        row=next(r for r in rows if r['source_locator']=='POR PROVINCIA!B3')
        self.assertIsNone(row['value']);self.assertEqual(row['value_raw'],'=SUM(1,2)')
        self.assertEqual(row['value_status'],'missing')
        self.assertIn('formula_result_not_cached',row['quality_flags'])

    def test_unexpected_or_missing_schema_is_rejected(self):
        for mode in ['sheet','column','category','missing_category']:
            with self.subTest(mode=mode):
                book=slaughter()
                if mode=='sheet':book.create_sheet('UNKNOWN')
                if mode=='column':book['POR PROVINCIA']['E2']='Unknown metric'
                if mode=='category':book['POR CATEGORÍA Y EDAD']['B2']='UNKNOWN'
                if mode=='missing_category':book['POR CATEGORÍA Y EDAD']['B2']=None
                with self.assertRaises(DataError):parse_workbook(encoded(book),'slaughter_detail',ASOF,2026)

    def test_invalid_values_future_date_wrong_year_or_cumulative_year_rejected(self):
        for bad in [-1,True,'NaN','not-a-number']:
            with self.subTest(value=bad):
                book=slaughter();book['POR PROVINCIA']['B3']=bad
                with self.assertRaises(DataError):parse_workbook(encoded(book),'slaughter_detail',ASOF,2026)
        with self.assertRaises(DataError):parse_workbook(encoded(slaughter()),'slaughter_detail',date(2026,1,31),2026)
        with self.assertRaises(DataError):parse_workbook(encoded(slaughter()),'slaughter_detail',ASOF,2025)
        book=slaughter(2019);book['POR ESTABLECIMIENTO']['G2']='ACUMULADO 2020'
        with self.assertRaises(DataError):parse_workbook(encoded(book),'slaughter_detail',ASOF,2019)

    def test_invalid_payload_kind_and_incomplete_families_rejected(self):
        with self.assertRaises(DataError):parse_workbook(b'not xlsx','slaughter_detail',ASOF)
        with self.assertRaises(DataError):parse_workbook(encoded(slaughter()),'unknown',ASOF)
        book=slaughter();book.remove(book['POR PROVINCIA'])
        with self.assertRaises(DataError):parse_workbook(encoded(book),'slaughter_detail',ASOF)
        book=production();book.active['A19']='AÑO 2026 - UNKNOWN'
        with self.assertRaises(DataError):parse_workbook(encoded(book),'production_detail',ASOF)

    def test_verified_2019_aggregate_conflict_is_quarantined_without_relabelling(self):
        book=slaughter(2019);sheet=book['POR ESTABLECIMIENTO'];sheet.delete_rows(1,sheet.max_row)
        sheet['A1']='AÑO 2019 - FAENA DE BOVINOS POR ESTABLECIMIENTO FAENADOR - EN CABEZAS'
        headers=['Nº Est.','CUIT','Razón Social Frigorifico','Provincia']+[
            'ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO','SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE','ACUMULADO 2020']
        for column,value in enumerate(headers,1):sheet.cell(3,column,value)
        for column,value in enumerate([69,30540080298,'Frigorífico','CORRIENTES']+[1]*12+[12],1):sheet.cell(4,column,value)
        rows=parse_workbook(encoded(book),'slaughter_detail',ASOF,2019)
        establishment=[row for row in rows if row['entity_type']=='establishment']
        self.assertEqual(len(establishment),12)
        self.assertEqual({row['frequency'] for row in establishment},{'monthly'})
        self.assertTrue(all(row['observation_date'].startswith('2019') for row in rows))
        flag='quarantined_aggregate_header:POR ESTABLECIMIENTO!Q3:ACUMULADO 2020'
        self.assertTrue(all(flag in row['quality_flags'] for row in rows))
        self.assertFalse(any(row['source_locator']=='POR ESTABLECIMIENTO!Q4' for row in rows))

    def test_2021_total_year_header_requires_complete_matching_year(self):
        months=['ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO','SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE']
        book=slaughter(2021);sheet=book['POR ESTABLECIMIENTO']
        for column,value in enumerate(months+['TOTAL AÑO 2021'],5):sheet.cell(2,column,value)
        for column in range(5,18):sheet.cell(3,column,12 if column==17 else 1)
        rows=parse_workbook(encoded(book),'slaughter_detail',ASOF,2021)
        annual=next(r for r in rows if r['entity_type']=='establishment' and r['frequency']=='annual')
        self.assertEqual(annual['source_locator'],'POR ESTABLECIMIENTO!Q3')
        self.assertEqual(annual['observation_date'],'2021-01-01')
        self.assertEqual(annual['period_end'],'2021-12-31')
        self.assertIn('published_annual',annual['quality_flags'])
        for text in ['TOTAL AÑO 2021','TOTAL AÑO 2020','TOTAL 2021']:
            book=slaughter(2021);book['POR ESTABLECIMIENTO']['G2']=text
            with self.subTest(header=text),self.assertRaises(DataError):
                parse_workbook(encoded(book),'slaughter_detail',ASOF,2021)

    def test_verified_2020_conflicting_entity_groups_quarantined_not_summed(self):
        book=slaughter(2020);sheet=book['POR ESTABLECIMIENTO']
        months=['ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO','SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE']
        for column,value in enumerate(months+['ACUMULADO 2020'],5):sheet.cell(2,column,value)
        defects=[(216,[422,30716396696,'CARNES ESPECIALES ENTRERRIANAS SRL','ENTRE RIOS'],155,7260),
                 (270,[422,30716396696,'CARNES ESPECIALES ENTRERRIANAS SRL','ENTRE RIOS'],841,2485),
                 (219,[47,30607813295,'LE GURIE S.A.','BUENOS AIRES'],1843,7208),
                 (226,[47,30607813295,'LE GURIE S.A.','BUENOS AIRES'],None,6694)]
        for row,metadata,october,total in defects:
            for column,value in enumerate(metadata,1):sheet.cell(row,column,value)
            sheet.cell(row,14,october);sheet.cell(row,17,total)
        rows=parse_workbook(encoded(book),'slaughter_detail',ASOF,2020)
        self.assertFalse(any(row['entity_id'] in {'422','47'} for row in rows))
        self.assertTrue(any(row['entity_id']=='69' for row in rows))
        self.assertTrue(any(row['entity_type']=='province' for row in rows))
        self.assertTrue(any(row['entity_type']=='category' for row in rows))
        flags={flag for row in rows for flag in row['quality_flags'] if flag.startswith('quarantined_conflicting_entity_rows:')}
        self.assertEqual(flags,{
            'quarantined_conflicting_entity_rows:POR ESTABLECIMIENTO!A216:Q216,A270:Q270:entity_id=422',
            'quarantined_conflicting_entity_rows:POR ESTABLECIMIENTO!A219:Q219,A226:Q226:entity_id=47'})
        self.assertTrue(all(flags.issubset(row['quality_flags']) for row in rows))


if __name__=='__main__':unittest.main()
