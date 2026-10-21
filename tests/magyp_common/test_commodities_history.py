import csv
import gzip
import io
import json
import sqlite3
import unittest
from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from openpyxl import Workbook
from scripts.magyp.grains import references as r, history_acquire as a, history_model as h, history_build as b
from scripts.magyp.build_price_freshness import add_commodities_history

ASOF=date(2026,10,21)


def manifest(family,**extra):
    return dict(family=family,source_url=r.SOURCES[family],capture_id='test-capture',
        sha256='a'*64,captured_at_utc='2026-10-21T12:00:00+00:00',http_status=200,**extra)


def xlsx(workbook):
    stream=io.BytesIO();workbook.save(stream);workbook.close();return stream.getvalue()


def board(day='2018-01-02'):
    return {'text':'Cotizaciones provisorias martes 02 enero 2018', 'rows':[
        ['PRODUCTO','BUENOS AIRES','BAHIA BLANCA','QUEQUEN','ROSARIO','CORDOBA']]+[
        [name,'1.234,50','0,00','-','100,25','0,00'] for name in ['GIRASOL','TRIGO PAN','MAIZ','SOJA','SORGO','CEBADA FORRAJERA']]}


class HistoryModelTests(unittest.TestCase):
    def test_internal_workbook_keeps_seventh_product_lino_and_original_month(self):
        w=Workbook();w.remove(w.active)
        headings=[['DARSENA','ROSARIO','B.BLANCA','CÓRDOBA','QUEQUEN','DARSENA','ROSARIO','B.BLANCA','QUEQUEN','DARSENA','ROSARIO','B.BLANCA','B.BLANCA','QUEQUEN'],
                  ['DARSENA','ROSARIO','B.BLANCA','QUEQUEN','DARSENA','ROSARIO','B.BLANCA','QUEQUEN','DARSENA','ROSARIO','QUEQUEN']]
        for name,markets in zip(['predar t m ysorg','predar s g y lino'],headings):
            s=w.create_sheet(name);s.cell(2,3,'EVOLUCION DE PRECIOS INTERNOS DE LOS PRINCIPALES GRANOS ( $/tn.)')
            for col,market in enumerate(markets,3):s.cell(6,col,market)
            s.cell(7,2,1996);s.cell(8,2,'ENERO')
            for col in range(3,3+len(markets)):s.cell(8,col,123.5)
        rows=list(h.internal_xlsx(xlsx(w),manifest('internal'),ASOF))
        self.assertEqual(len(rows),25);self.assertEqual(sum(x['commodity']=='LINO' for x in rows),3)
        self.assertTrue(all(x['fecha']=='1996-01-01' and x['frecuencia']=='mensual' for x in rows))

    def test_internal_annual_lino_is_not_relabelled_as_barley(self):
        obj={'text':'PRECIOS INTERNOS DE LOS PRINCIPALES GRANOS ( $/tn.)','rows':[
            ['AÑO \\ MES','TRIGO','MAIZ','SORGO'],
            ['DARSENA','ROSARIO','B.BLANCA','QUEQUEN','DARSENA','ROSARIO','B.BLANCA','QUEQUEN','DARSENA','ROSARIO','B.BLANCA'],
            ['2018'],['ENERO']+['100,00']*11,
            ['AÑO \\ MES','SOJA','GIRASOL','LINO'],
            ['DARSENA','ROSARIO','B.BLANCA','QUEQUEN']*2+['DARSENA','ROSARIO','QUEQUEN'],
            ['2018'],['ENERO']+['100,00']*11]}
        m=manifest('internal');m['source_url']='https://www.magyp.gob.ar/100002_2018.php'
        rows=list(h.internal_annual(obj,m,ASOF))
        self.assertEqual(len(rows),22);self.assertEqual(sum(x['commodity']=='LINO' for x in rows),3)
        self.assertFalse(any(x['commodity']=='CEBADA FORRAJERA' for x in rows))

    def test_board_original_values_zero_and_exact_date(self):
        m=manifest('board');m['source_url']='https://www.magyp.gob.ar/camaras2018-01-02.php'
        rows=list(h.board_archive(board(),m,ASOF))
        self.assertEqual(len(rows),30)
        self.assertEqual(rows[0]['price'],'1234.50');self.assertEqual(rows[0]['price_raw'],'1.234,50')
        self.assertEqual(rows[1]['value_status'],'zero');self.assertEqual(rows[2]['value_status'],'missing')
        self.assertEqual(rows[0]['raw_sha256'],m['sha256'])
        m['source_url']=m['source_url'].replace('01-02','01-03')
        with self.assertRaises(r.Error):list(h.board_archive(board(),m,ASOF))

    def test_board_rejects_unknown_product_or_column(self):
        m=manifest('board');m['source_url']='https://www.magyp.gob.ar/camaras2018-01-02.php'
        for mutation in ('product','column'):
            obj=board()
            if mutation=='product':obj['rows'][1][0]='UNKNOWN'
            else:obj['rows'][1].append('99')
            with self.assertRaises(r.Error):list(h.board_archive(obj,m,ASOF))

    def test_futures_date_is_not_maturity_and_two_digit_year_is_anchored(self):
        m=manifest('futures');m['source_url']='https://www.magyp.gob.ar/_chicago/12_maiz_1995.php'
        obj={'text':'MERCADO DE CHICAGO - PROMEDIOS MENSUALES EN DOL/TON MAIZ Enero',
             'rows':[['INICIO EMBARQ','03/01/95','04/01/1995','Promedio'],['MAY1996','89,96','90.25','90,10']]}
        rows=list(h.futures_matrix(obj,m,ASOF))
        self.assertEqual([x['fecha'] for x in rows],['1995-01-03','1995-01-04'])
        self.assertTrue(all(x['shipment_window']=='1996-05' for x in rows))
        self.assertTrue(all(x['position_raw']=='MAY1996' for x in rows))
        self.assertEqual(rows[1]['price'],'90.25');self.assertEqual(rows[0]['tipo_precio'],'Futuro / histórico publicado')
        obj['rows'][0][1]='03/01/96'
        with self.assertRaises(r.Error):list(h.futures_matrix(obj,m,ASOF))

    def test_futures_invalid_actual_date_is_not_repaired(self):
        m=manifest('futures');m['source_url']='https://www.magyp.gob.ar/_chicago/12_maiz_1995.php'
        obj={'text':'MERCADO DE CHICAGO - PROMEDIOS MENSUALES EN DOL/TON MAIZ Enero',
             'rows':[['INICIO EMBARQ','31/02/1995','Promedio'],['MAY1995','12,50','12,50']]}
        with self.assertRaises(ValueError):list(h.futures_matrix(obj,m,ASOF))

    def test_futures_identical_duplicate_collapses_but_conflict_rejects(self):
        m=manifest('futures');m['source_url']='https://www.magyp.gob.ar/_chicago/12_maiz_1995.php'
        obj={'text':'MERCADO DE CHICAGO - PROMEDIOS MENSUALES EN DOL/TON MAIZ Enero',
             'rows':[['INICIO EMBARQ','03/01/1995','Promedio'],['MAY1995','12,50','12,50'],['MAY1995','12,50','12,50']]}
        self.assertEqual(len(list(h.futures_matrix(obj,m,ASOF))),1)
        obj['rows'][-1][1]='13,50'
        with self.assertRaises(r.Error):list(h.futures_matrix(obj,m,ASOF))

    def test_old_daily_futures_context_is_verified_without_relaxing_current_parser(self):
        obj={'text':'5 de enero de 2021 COTIZACIONES DE CIERRE * En Dolares Estadounidenses/Tn.',
             'rows':[['CHICAGO *'],['POSICION','TRIGO','MAIZ','AVENA','SOJA','HARINA DE SOJA','ACEITE DE SOJA'],
                ['5/ene','4/ene']*6,['MAY2021']+['100,25','99,75']*6,
                ['KANSAS *'],['POSICION','TRIGO'],['5/ene','4/ene'],['MAY2021','100,25','99,75'],
                ['COTIZACIONES FOB BS. AS.'],['MAY2021']+['999,99']*12]}
        m=manifest('futures');m['source_url']='https://www.magyp.gob.ar/cotiza_2021-01-05.php'
        with self.assertRaises(r.Error):r.parse_futures(obj)
        rows=list(h.futures_daily(obj,m,ASOF))
        self.assertEqual(len(rows),14);self.assertEqual(rows[0]['price_raw'],'100,25')
        self.assertTrue(all(row['tipo_precio']=='Futuro / cierre' for row in rows))
        self.assertIn('historical_context_spelling_variant',rows[0]['quality_flags'])
        obj['text']=obj['text'].replace('Dolares Estadounidenses','moneda desconocida')
        with self.assertRaises(r.Error):list(h.futures_daily(obj,m,ASOF))

    def test_fas_both_verified_header_offsets_keep_unknown_oil_units(self):
        for header in (6,7):
            w=Workbook();s=w.active;s.cell(1,1,'FAS - Derechos de Exportacion Completos')
            for col,value in enumerate(['Fecha']+r.FAS_LABELS,1):s.cell(header,col,value)
            for col,value in enumerate([datetime(2007,1,16)]+[100]*9,1):s.cell(header+1,col,value)
            m=manifest('fas',regime='D.E.C.',requested_from='2007-01-01',requested_to='2026-10-21')
            rows=list(h.fas_xlsx(xlsx(w),m,ASOF))
            self.assertEqual(len(rows),9);self.assertEqual(rows[0]['moneda'],'ARS')
            self.assertEqual(rows[-1]['moneda'],'Sin identificar');self.assertEqual(rows[-1]['unidad'],'Sin identificar')
            self.assertIn('sunflower_formula_pre_2019_02_14',rows[6]['quality_flags'])

    def test_fob_catalog_exact_schema_and_calendar_index_preserved(self):
        d={'field':[{'title':'indice_tiempo'}, {'title':'trigo','description':'Trigo publicado','id':'official-id','units':'Dolares por tonelada'}]}
        body=b'indice_tiempo,trigo\n1993-01-09,100.5\n1993-01-10,100.5\n'
        rows=list(h.fob_sspm(body,manifest('fob'),d,ASOF))
        self.assertEqual(len(rows),2);self.assertEqual(rows[0]['fecha'],'1993-01-09')
        self.assertIn('weekend_value_published_by_sspm',rows[0]['quality_flags'])
        self.assertEqual(rows[0]['product_code'],'official-id');self.assertFalse(rows[0].get('position_raw'))
        self.assertEqual(rows[0]['moneda'],'USD')
        with self.assertRaises(r.Error):list(h.fob_sspm(body.replace(b'trigo',b'unknown'),manifest('fob'),d,ASOF))

    def test_discovery_only_published_links_and_public_official_urls(self):
        body=b"<a href=\"javascript:window.open('12_maiz_1995.php','x')\">1995</a><a href='https://evil.test/data'>x</a>"
        links=a.published_links(body,'https://www.magyp.gob.ar/index.php')
        self.assertEqual(links,['https://www.magyp.gob.ar/12_maiz_1995.php'])
        for url in ('https://user:secret@www.magyp.gob.ar/a','https://www.magyp.gob.ar/a?token=secret'):
            with self.assertRaises(r.Error):a.official_url(url)

    def test_missing_markers_not_zero_or_bool(self):
        self.assertEqual(h.numeric('s/c'),('','missing'))
        self.assertEqual(h.numeric('--'),('','missing'))
        self.assertEqual(h.numeric(0),('0','zero'))
        with self.assertRaises(r.Error):h.numeric(True)


class HistoryBuildTests(unittest.TestCase):
    def test_partial_file_rolls_back_and_hashes_validate(self):
        with TemporaryDirectory() as directory:
            root=Path(directory);good=manifest('board');bad={**good,'capture_id':'bad'}
            row=h.record('board','2018-01-02','SOJA','ROSARIO','Pizarra provisional','ARS','TN','100',good,'row:1')
            def parser(path,m,as_of):
                yield {**row,'capture_id':m['capture_id']}
                if m['capture_id']=='bad':raise r.Error('Unknown final block')
            with patch.object(b,'captures',return_value=[(None,good),(None,bad)]),patch.object(b,'parse_capture',side_effect=parser):
                entry=b.build_family(root,'board',ASOF)
            self.assertEqual(entry['coverage']['positive_rows'],1)
            self.assertEqual(len(entry['coverage']['rejected_files']),1)
            db=sqlite3.connect(root/'analytical/commodities/history/board.sqlite')
            self.assertEqual(json.loads(db.execute('select payload from obs').fetchone()[0])['capture_id'],'test-capture');db.close()
            path=root/'dashboard/commodities/history/board'/entry['monthly']['path']
            self.assertEqual(r.digest(path.read_bytes()),entry['monthly']['sha256'])
            self.assertEqual(len(list(csv.DictReader(io.StringIO(gzip.decompress(path.read_bytes()).decode()),delimiter=';'))),1)

    def test_no_valid_capture_does_not_overwrite_published_index(self):
        with TemporaryDirectory() as directory:
            root=Path(directory);target=root/'dashboard/commodities/history/index.json'
            target.parent.mkdir(parents=True);target.write_bytes(b'previous-valid-index')
            with patch.object(b,'captures',return_value=[]):
                with self.assertRaises(r.Error):b.build_family(root,'board',ASOF)
            self.assertEqual(target.read_bytes(),b'previous-valid-index')

    def test_invalid_revision_of_previously_valid_file_preserves_index(self):
        with TemporaryDirectory() as directory:
            root=Path(directory);old=manifest('board')
            row=h.record('board','2018-01-02','SOJA','ROSARIO','Pizarra provisional','ARS','TN','100',old,'row:1')
            with patch.object(b,'captures',return_value=[(None,old)]),patch.object(b,'parse_capture',return_value=iter([row])):
                entry=b.build_family(root,'board',ASOF)
            target=root/'dashboard/commodities/history/index.json'
            target.write_bytes(r.encoded({'schema_version':h.VERSION,'families':{'board':entry}}));before=target.read_bytes()
            bad={**old,'capture_id':'bad-revision'}
            new={**old,'capture_id':'new-resource','source_url':'https://www.magyp.gob.ar/another.php'}
            def parser(path,m,as_of):
                if m['capture_id']=='bad-revision':raise r.Error('Unknown changed schema')
                yield {**row,'capture_id':m['capture_id']}
            with patch.object(b,'captures',return_value=[(None,bad),(None,new)]),patch.object(b,'parse_capture',side_effect=parser):
                with self.assertRaisesRegex(r.Error,'antes válido'):b.build_family(root,'board',ASOF)
            self.assertEqual(target.read_bytes(),before)

    def test_raw_hash_mismatch_rejected(self):
        with TemporaryDirectory() as directory:
            p=Path(directory)/'manifest.json';(p.parent/'response.json').write_text('{}')
            m={**manifest('board'),'raw_file':'response.json','size_bytes':2}
            with self.assertRaises(r.Error):list(b.parse_capture(p,m,ASOF))

    def test_fas_variation_not_across_formula_change(self):
        self.assertNotEqual(b.method_key({'quality_flags':['sunflower_formula_pre_2019_02_14']}),
                            b.method_key({'quality_flags':['sunflower_formula_from_2019_02_14']}))

    def test_freshness_enrichment_does_not_touch_other_modules(self):
        with TemporaryDirectory() as directory:
            root=Path(directory);p=root/'data/magyp/dashboard/commodities/history/index.json';p.parent.mkdir(parents=True)
            p.write_text(json.dumps({'schema_version':h.VERSION,'families':{}}))
            sources=[{'source_family':key,'status':'STATIC','custom':'untouched'} for key in ('mcba','corrientes','sio')]
            self.assertEqual(add_commodities_history({'sources':sources},root),{'sources':sources})

    def test_historical_success_never_masks_current_fob_pipeline_error(self):
        with TemporaryDirectory() as directory:
            root=Path(directory);p=root/'data/magyp/dashboard/commodities/history/index.json';p.parent.mkdir(parents=True)
            coverage=dict(oldest_verified_date='1993-01-04',latest_verified_date='2025-01-21',positive_rows=10,
                monthly_rows=2,products=['Trigo'],markets=['SSPM'],currency=['USD'],unit=['TN'],
                coverage_status='VERIFIED_WITH_GAPS',last_successful_capture='2026-10-21T12:00:00Z')
            p.write_text(json.dumps({'schema_version':h.VERSION,'families':{'fob':{'coverage':coverage}}}))
            for status in ('PIPELINE_ERROR','SOURCE_UNAVAILABLE','NO_PUBLICATION'):
                row=dict(source_family='magyp_fob',status=status,last_successful_update='2026-10-19T12:00:00Z',
                    last_observation_date='2026-10-19',evidence='current API workflow',attempted_at='2026-10-20T12:00:00Z')
                result=add_commodities_history({'sources':[row.copy()]},root)['sources'][0]
                for key in row:self.assertEqual(result[key],row[key])
                self.assertEqual(result['historical_coverage']['status'],'VALIDATED_BACKFILL')


if __name__=='__main__':unittest.main()
